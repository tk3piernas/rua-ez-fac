"""Writing to DEUDA A PAGAR ARG.xlsx (stage 4).

The logic (which row, what goes in each column, what to journal, how to undo) is written once
against a small adapter. Two adapters exist:
  * LibroCOM      desktop Excel through pywin32: the only one used on real files. Works whether
                  the workbook is already open (it writes into that same Excel) or closed (it
                  opens it in a hidden Excel, saves and closes).
  * LibroOpenpyxl only for tests in the cloud, where there is no Excel.
"""
from __future__ import annotations

import re
import shutil
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from urllib.parse import quote, unquote

PRIMERA, ULTIMA = 12, 210
COLS = "ABCDEFGHIJ"
COLS_PROV = "ABCDEFGHIJKLMNO"
AMARILLO = 13431551                      # light yellow (same as the PS1), BGR integer
TEXTO_PROV = ("C", "G", "K", "L")        # CUIT, phone, account number, CBU: keep as text


class ExcelError(Exception):
    pass


RPC_E_CALL_REJECTED = -2147418111        # Excel busy: starting, or a hidden dialog waiting for an answer
EXCEL_NO_RESPONDE = ("Excel no responde (puede estar mostrando un cartel escondido). Cerrá los \"Microsoft Excel\" "
                     "del Administrador de tareas, abrí Excel una vez a mano para ver si pide algo, y volvé a probar")


def _ocupado(e: Exception) -> bool:
    return bool(getattr(e, "args", None)) and e.args[0] == RPC_E_CALL_REJECTED


def _pid_de(xl) -> int | None:
    """Process id of a hidden Excel started with DispatchEx (from its window), or None."""
    try:
        import win32process
        return win32process.GetWindowThreadProcessId(_reintentar_ocupado(lambda: xl.Hwnd, segundos=5))[1] or None
    except Exception:
        return None


def _terminar_excel(pid: int) -> bool:
    """Ends that process, only if it really is EXCEL.EXE (a reused pid is left alone)."""
    import ctypes
    from ctypes import wintypes
    k32 = ctypes.WinDLL("kernel32", use_last_error=True)
    k32.OpenProcess.restype = wintypes.HANDLE
    h = k32.OpenProcess(0x0001 | 0x1000, False, pid)        # PROCESS_TERMINATE | PROCESS_QUERY_LIMITED_INFORMATION
    if not h:
        return False
    try:
        buf, n = ctypes.create_unicode_buffer(1024), wintypes.DWORD(1024)
        if not k32.QueryFullProcessImageNameW(h, 0, buf, ctypes.byref(n)) or \
                not buf.value.lower().endswith("\\excel.exe"):
            return False
        return bool(k32.TerminateProcess(h, 1))
    finally:
        k32.CloseHandle(h)


def _reintentar_ocupado(fn, segundos: float = 20, pausa: float = 1):
    """`fn()`, retried while Excel answers 'call was rejected' (it does that for a while when it starts)."""
    import time
    fin = time.monotonic() + segundos
    while True:
        try:
            return fn()
        except Exception as e:
            if not _ocupado(e) or time.monotonic() >= fin:
                raise
            time.sleep(pausa)


class LibroConCambios(ExcelError):
    """The workbook is open with unsaved changes: the app never saves the user's work on its own."""


def serial(d: date) -> float:
    return float((d - date(1899, 12, 30)).days)


def _col(c: str) -> int:
    return ord(c) - 64


# ======================================================================= adapters

class LibroOpenpyxl:
    """Test adapter (no Excel)."""

    def __init__(self, ruta: Path):
        import openpyxl
        self.ruta = ruta
        self.wb = openpyxl.load_workbook(ruta)
        self.ya_abierto = False
        self.saved = True

    def hoja(self, nombre):
        if nombre not in self.wb.sheetnames:
            raise ExcelError(f"la planilla no tiene la hoja {nombre}")
        return HojaOpenpyxl(self.wb[nombre])

    def nombres(self) -> list[str]:
        return list(self.wb.sheetnames)

    def copiar_hoja(self, modelo: str, nombre: str):
        """New sheet `nombre` right after `modelo`, with its layout (formats, formulas, widths)."""
        ws = self.wb.copy_worksheet(self.wb[modelo])
        ws.title = nombre
        self.wb.move_sheet(ws, self.wb.sheetnames.index(modelo) + 1 - self.wb.sheetnames.index(nombre))
        return HojaOpenpyxl(ws)

    def guardar(self):
        self.wb.save(self.ruta)

    def marcar_guardado(self):
        self.saved = True

    def guardar_copia(self, destino: Path):
        shutil.copy2(self.ruta, destino)

    def cerrar(self):
        self.wb.close()


class HojaOpenpyxl:
    def __init__(self, ws):
        self.ws = ws

    def formula(self, r, c):
        v = self.ws[f"{c}{r}"].value
        if v is None:
            return ""
        if isinstance(v, datetime):
            return serial(v.date())
        return v

    def valor(self, r, c):
        return self.formula(r, c) if not str(self.formula(r, c)).startswith("=") else None

    def poner(self, r, c, v, texto=False, formula=False):
        """`formula`: a value starting with "=" is a formula (only H and J of the rows, and undo); otherwise
        it is written as text, so a pasted "=ver mail" never becomes #NAME?."""
        cell = self.ws[f"{c}{r}"]
        if texto:
            cell.number_format = "@"
        cell.value = v if v != "" else None
        if isinstance(v, str) and v.startswith("=") and not formula:
            cell.number_format = "@"
            cell.data_type = "s"

    def color(self, r, c):
        f = self.ws[f"{c}{r}"].fill
        return None if not f or f.fill_type is None else f.fgColor.rgb

    def pintar(self, r, c, color):
        from openpyxl.styles import PatternFill
        self.ws[f"{c}{r}"].fill = PatternFill() if color is None else PatternFill("solid", fgColor=color if isinstance(color, str) else "FFFFF2CC")

    def comentario(self, r, c):
        cm = self.ws[f"{c}{r}"].comment
        return cm.text if cm else None

    def comentar(self, r, c, texto):
        from openpyxl.comments import Comment
        self.ws[f"{c}{r}"].comment = Comment(texto, "RUA EZ-Fac") if texto else None

    def limpiar(self, rango):
        for fila in self.ws[rango]:
            for cell in fila:
                cell.value = None
                cell.hyperlink = None
                cell.comment = None

    def lista(self, rango, formula):
        """Dropdown (data validation) for `rango`, replacing the one copied from the model sheet."""
        from openpyxl.worksheet.datavalidation import DataValidation
        self.ws.data_validations.dataValidation = [d for d in self.ws.data_validations.dataValidation
                                                   if str(d.sqref) != rango]
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        self.ws.add_data_validation(dv)
        dv.add(rango)

    def vaciar_rango(self, rango, sin_columna="I", formula_h_desde=None):
        """Data rows back to empty: values, fills, comments and links; H gets its formula back; I untouched."""
        from openpyxl.formula.translate import Translator
        from openpyxl.styles import PatternFill
        ref = self.ws[f"H{formula_h_desde}"].value if formula_h_desde else None
        for fila in self.ws[rango]:
            for cell in fila:
                cell.fill = PatternFill()
                cell.comment = None
                if cell.column_letter == sin_columna:
                    continue
                cell.hyperlink = None
                if cell.column_letter == "H":
                    cell.value = (Translator(ref, origin=f"H{formula_h_desde}").translate_formula(cell.coordinate)
                                  if ref else f'=IF(G{cell.row}="","",ROUND(G{cell.row}*$B$4,2))')
                else:
                    cell.value = None


class LibroCOM:
    """Desktop Excel. Finds the workbook if it is already open (Running Object Table), else opens it hidden."""

    def __init__(self, ruta: Path, url_onedrive: str = ""):
        import pythoncom
        import win32com.client
        pythoncom.CoInitialize()
        self.ruta = ruta
        self.xl = None
        self.nuevo = False
        self.pid = None                    # process of the hidden Excel this app opened (never the user's)
        self._hojas: list = []
        self.wb = self._buscar_abierto(ruta, url_onedrive)
        self.ya_abierto = self.wb is not None
        if not self.wb:
            try:
                self.xl = win32com.client.DispatchEx("Excel.Application")
                self.nuevo = True
                self.pid = _pid_de(self.xl)
                self.xl.Visible = False
                self.xl.DisplayAlerts = False
                self.wb = _reintentar_ocupado(lambda: self.xl.Workbooks.Open(str(ruta), 0, False))
            except Exception as e:
                self.cerrar()
                if _ocupado(e):
                    raise ExcelError(EXCEL_NO_RESPONDE) from e
                raise ExcelError(f"no pude abrir Excel: {e}") from e
        if self.wb.ReadOnly:
            self.cerrar()
            raise ExcelError("la planilla está abierta en modo solo lectura (¿la tiene abierta otra persona?)")

    @staticmethod
    def _buscar_abierto(ruta: Path, url_onedrive: str):
        """The workbook in the Running Object Table whose name is exactly `ruta` or `url_onedrive`."""
        import pythoncom
        import win32com.client
        rot = pythoncom.GetRunningObjectTable()
        ctx = pythoncom.CreateBindCtx(0)
        for m in rot:
            try:
                nombre = m.GetDisplayName(ctx, None)
            except Exception:
                continue
            if not es_el_libro(nombre, ruta, url_onedrive):
                continue
            try:
                obj = rot.GetObject(m)
                wb = win32com.client.Dispatch(obj.QueryInterface(pythoncom.IID_IDispatch))
                if es_el_libro(str(wb.FullName), ruta, url_onedrive):
                    return wb
            except Exception:
                continue
        return None

    @property
    def saved(self):
        return bool(self.wb.Saved)

    def hoja(self, nombre):
        try:
            h = HojaCOM(self.wb.Worksheets(nombre))
        except Exception as e:
            raise ExcelError(f"la planilla no tiene la hoja {nombre}") from e
        self._hojas.append(h)
        return h

    def nombres(self) -> list[str]:
        return [self.wb.Worksheets(i).Name for i in range(1, self.wb.Worksheets.Count + 1)]

    def copiar_hoja(self, modelo: str, nombre: str):
        """New sheet `nombre` right after `modelo`, with its layout (formats, formulas, widths)."""
        m = self.wb.Worksheets(modelo)
        m.Copy(None, m)                       # Copy(Before, After)
        ws = self.wb.Worksheets(m.Index + 1)
        ws.Name = nombre
        h = HojaCOM(ws)
        self._hojas.append(h)
        return h

    def guardar(self):
        self.wb.Save()

    def marcar_guardado(self):
        """After a revert the cells are as they were on disk: tell Excel there is nothing to save."""
        self.wb.Saved = True

    def guardar_copia(self, destino: Path):
        """Copy with what Excel has now (also unsaved changes); the open workbook stays as it is."""
        self.wb.SaveCopyAs(str(destino))

    def cerrar(self):
        import gc
        try:
            for h in getattr(self, "_hojas", []):
                h.ws = None                   # COM sheet objects released before Excel quits
            if self.wb is not None and not self.ya_abierto:
                self.wb.Close(False)
            self.wb = None                    # released while Excel is still running
            gc.collect()
            if self.xl is not None and self.nuevo:
                self.xl.Quit()
        except Exception:
            # a hung hidden Excel would stay in the background and make the next ones reject every call
            if self.nuevo and self.pid:
                _terminar_excel(self.pid)
        self.wb = self.xl = None


class HojaCOM:
    def __init__(self, ws):
        self.ws = ws

    def _c(self, r, c):
        return self.ws.Range(f"{c}{r}")

    def formula(self, r, c):
        f = self._c(r, c).Formula
        if isinstance(f, float) and f.is_integer():
            return f
        return f if f is not None else ""

    def valor(self, r, c):
        return self._c(r, c).Value2

    def poner(self, r, c, v, texto=False, formula=False):
        """`formula`: a value starting with "=" is a formula (only H and J of the rows, and undo); otherwise
        it is written as text, so a pasted "=ver mail" never becomes #NAME?."""
        cell = self._c(r, c)
        if texto:
            cell.NumberFormat = "@"
        if isinstance(v, str) and v.startswith("=") and not formula:
            cell.NumberFormat = "@"
            cell.Value2 = v
        elif isinstance(v, str) and v.startswith("="):
            cell.Formula = v
        elif v == "" or v is None:
            cell.ClearContents()
        else:
            cell.Value2 = v

    def color(self, r, c):
        it = self._c(r, c).Interior
        return None if it.ColorIndex == -4142 else int(it.Color)

    def pintar(self, r, c, color):
        it = self._c(r, c).Interior
        if color is None:
            it.ColorIndex = -4142
        else:
            it.Color = AMARILLO if color is True or isinstance(color, str) else color

    def comentario(self, r, c):
        cm = self._c(r, c).Comment
        return cm.Text() if cm is not None else None

    def comentar(self, r, c, texto):
        cell = self._c(r, c)
        if cell.Comment is not None:
            cell.Comment.Delete()
        if texto:
            cell.AddComment(texto)

    def limpiar(self, rango):
        rg = self.ws.Range(rango)
        rg.Hyperlinks.Delete()
        rg.ClearContents()
        rg.ClearComments()

    def lista(self, rango, formula):
        """Dropdown (data validation) for `rango`, replacing the one copied from the model sheet."""
        v = self.ws.Range(rango).Validation
        v.Delete()
        v.Add(3, 1, 1, formula)               # xlValidateList, xlValidAlertStop, xlBetween

    def vaciar_rango(self, rango, sin_columna="I", formula_h_desde=None):
        """Data rows back to empty: values, fills, comments and links; H gets its formula back; I untouched."""
        primera, ultima = (int("".join(ch for ch in x if ch.isdigit())) for x in rango.split(":"))
        r1c1 = self.ws.Range(f"H{formula_h_desde}").FormulaR1C1 if formula_h_desde else \
            '=IF(RC[-1]="","",ROUND(RC[-1]*R4C2,2))'
        todo = self.ws.Range(rango)
        todo.Interior.ColorIndex = -4142
        todo.ClearComments()
        self.ws.Range(f"J{primera}:J{ultima}").Hyperlinks.Delete()
        self.ws.Range(f"A{primera}:G{ultima}").ClearContents()
        self.ws.Range(f"J{primera}:J{ultima}").ClearContents()
        self.ws.Range(f"H{primera}:H{ultima}").FormulaR1C1 = r1c1


def abrir(cfg, ruta: Path | None = None):
    ruta = ruta or cfg.planilla
    if not ruta or not ruta.exists():
        raise ExcelError(f"no encuentro la planilla ({ruta})")
    if getattr(cfg, "excel_motor", "com") == "openpyxl":
        return LibroOpenpyxl(ruta)
    try:
        import win32com.client  # noqa: F401
    except ImportError as e:
        raise ExcelError("falta pywin32: corré 'Instalar o actualizar RUA EZ-Fac.bat'") from e
    return LibroCOM(ruta, url_onedrive=url_onedrive_de(cfg, ruta))


def _misma_ruta(a, b) -> bool:
    try:
        return Path(a).resolve() == Path(b).resolve()
    except (OSError, ValueError):
        return str(a).lower() == str(b).lower()


def es_el_libro(nombre: str, ruta: Path, url: str) -> bool:
    """`nombre` (a ROT display name or a workbook FullName) is this very workbook: same local path,
    or the full SharePoint URL of the real sheet. A workbook with the same file name elsewhere is not."""
    n = str(nombre or "")
    if n.lower().startswith("http"):
        # the ROT has it percent-encoded, FullName usually not: compare both decoded
        return bool(url) and unquote(n).rstrip("/").lower() == unquote(url).rstrip("/").lower()
    return _misma_ruta(n, ruta) if n else False


def url_onedrive_de(cfg, ruta: Path) -> str:
    """Full SharePoint URL of `ruta` only when it is inside the real synced OneDrive folder of
    config.toml (the real sheet, PROVEEDORES RUA.xlsx...); "" for the test copy or anything else,
    so an open workbook is then matched only by its local path."""
    from . import config
    try:
        real = config.cargar(cfg.archivo) if getattr(cfg, "archivo", None) else config.cargar()
    except Exception:
        return ""
    if not real.base or not real.url_base:
        return ""
    try:
        rel = Path(ruta).resolve().relative_to(Path(real.base).resolve().parent)
    except (OSError, ValueError):
        return ""
    raiz = real.url_base.rstrip("/").rsplit("/", 1)[0]       # url_base is 'Facturas pendientes'
    return raiz + "/" + "/".join(quote(p) for p in rel.parts)


# ======================================================================= writing

def _vacia(h, r) -> bool:
    return all(h.formula(r, c) in ("", None) for c in "ABCDEFGJ")


def _fila_libre(h, ocupadas: set) -> int | None:
    for r in range(PRIMERA, ULTIMA + 1):
        if r not in ocupadas and _vacia(h, r):
            return r
    return None


def _hyperlink(url: str, texto: str) -> str:
    q = lambda s: s.replace('"', '""')
    return f'=HYPERLINK("{q(url)}","{q(texto)}")'


RESPALDOS_A_GUARDAR = 20


def respaldar(ruta: Path, carpeta: Path, motivo: str = "") -> Path:
    """Copy of the workbook in `carpeta` (datos\\backups); only the newest RESPALDOS_A_GUARDAR are kept."""
    carpeta.mkdir(parents=True, exist_ok=True)
    dest = carpeta / f"{datetime.now():%Y-%m-%d_%H%M%S} {motivo + ' ' if motivo else ''}{ruta.name}"
    shutil.copy2(ruta, dest)
    podar_respaldos(carpeta)
    return dest


def podar_respaldos(carpeta: Path, quedan: int = RESPALDOS_A_GUARDAR) -> list[Path]:
    """Deletes the oldest workbook copies (their names start with the date) beyond `quedan`."""
    copias = sorted(carpeta.glob("*.xlsx"), key=lambda p: p.name)
    viejas = copias[:-quedan] if len(copias) > quedan else []
    for v in viejas:
        try:
            v.unlink()
        except OSError:
            pass
    return viejas


def escribir(cfg, lote, diario, backups: Path, guardar_cambios_usuario: bool = False, log=print) -> dict:
    """Writes the batch rows and the new providers. Returns {"filas": n, "proveedores": n}."""
    filas = [f for f in lote.filas if not f.excluida]
    provs = lote.proveedores
    if not filas and not provs:
        return {"filas": 0, "proveedores": 0}
    libro = abrir(cfg)
    try:
        if not libro.saved:
            if not guardar_cambios_usuario:
                raise LibroConCambios("la planilla está abierta con cambios sin guardar")
            libro.guardar()
        copia = respaldar(Path(cfg.planilla), backups)
        diario.anotar({"op": "respaldo_excel", "ruta": str(copia)})
        hojas = {}
        ocupadas: dict[str, set] = {}
        hechos: list[dict] = []
        anotar = diario.anotar

        def anotar_y_guardar(paso):
            hechos.append(paso)
            anotar(paso)
        diario_anotar = anotar_y_guardar
        try:
            n, np = _escribir_todo(cfg, libro, filas, provs, hojas, ocupadas, diario_anotar, log)
        except Exception:
            # nothing half-written stays in the workbook: undo this run in memory and do not save
            _revertir(libro, hechos)
            if libro.ya_abierto:
                # it was saved when we started (checked above), and the revert left it as on disk
                try:
                    libro.marcar_guardado()
                except Exception:
                    pass
            anotar({"op": "excel_revertido"})
            raise
        libro.guardar()
        diario.anotar({"op": "excel_guardado"})
        return {"filas": n, "proveedores": np}
    finally:
        libro.cerrar()


def _escribir_todo(cfg, libro, filas, provs, hojas, ocupadas, anotar, log):
    """Each step is journaled BEFORE its cells are written, so a failure halfway can be reverted."""
    n = 0
    for f in filas:
        h = hojas.setdefault(f.hoja, libro.hoja(f.hoja))
        usadas = ocupadas.setdefault(f.hoja, set())
        r = _fila_libre(h, usadas)
        if r is None:
            raise ExcelError(f"la hoja {f.hoja} está llena (fila {ULTIMA})")
        usadas.add(r)
        escrito: dict = {"A": f.proveedor}
        texto = set()
        if f.tarea:
            escrito["B"] = f.tarea
            texto.add("B")
        if cfg.escribir_servicio and f.servicio:
            escrito["C"] = f.servicio
        if f.empresa:
            escrito["D"] = f.empresa
        if f.fecha:
            escrito["E"] = serial(f.fecha)
        if f.vto:
            escrito["F"] = serial(f.vto)
        if f.monto is not None:
            escrito["G"] = float(f.monto)
        if f.iva is not None:
            escrito["H"] = float(f.iva)
        if f.url:
            escrito["J"] = _hyperlink(f.url, f.texto_link)
        anotar({"op": "fila_excel", "libro": str(cfg.planilla), "hoja": f.hoja, "fila": r,
                "antes": {c: h.formula(r, c) for c in COLS}, "escrito": escrito,
                "colores": {c: h.color(r, c) for c in COLS} if f.resaltar else {},
                "comentario_antes": h.comentario(r, "A") if f.comentario else None,
                "comentado": bool(f.comentario), "factura": f.factura_id})
        for c, v in escrito.items():
            h.poner(r, c, v, texto=c in texto, formula=c in ("H", "J"))
        if f.resaltar:
            for c in COLS:                        # I included on purpose: only its fill, never its formula (Matheo)
                h.pintar(r, c, True)
        if f.comentario:
            h.comentar(r, "A", f.comentario)
        if str(h.formula(r, "A")) != f.proveedor:
            raise ExcelError(f"verificación: la fila {r} de {f.hoja} no quedó bien")
        log(f"EXCEL {f.hoja} fila {r}: {f.proveedor} | {f.texto_link} | {f.monto}")
        n += 1
    np = 0
    if provs:
        h = libro.hoja("Proveedores")
        usadas = set()
        for p in provs:
            r = next((r for r in range(5, 301) if r not in usadas
                      and all(h.formula(r, c) in ("", None) for c in COLS_PROV)), None)
            if r is None:
                raise ExcelError("la hoja Proveedores está llena")
            usadas.add(r)
            anotar({"op": "proveedor_excel", "libro": str(cfg.planilla), "fila": r,
                    "valores": [str(v) for v in p["valores"]]})
            for c, v in zip(COLS_PROV, p["valores"]):
                if v not in ("", None):
                    h.poner(r, c, str(v), texto=c in TEXTO_PROV)
            log(f"EXCEL Proveedores fila {r}: {p['valores'][1]}")
            np += 1
    return n, np


def _restaurar(libro, p, problemas, log) -> None:
    if p["op"] == "fila_excel":
        h = libro.hoja(p["hoja"])
        r = p["fila"]
        if str(h.formula(r, "A")) != str(p["escrito"]["A"]) or (
                "J" in p["escrito"] and str(h.formula(r, "J")) != p["escrito"]["J"]):
            problemas.append(f"{p['hoja']} fila {r} cambió después del lote; no la toco")
            return
        for c in p["escrito"]:
            h.poner(r, c, p["antes"].get(c, ""), formula=True)      # back exactly as it was
        for c, col in (p.get("colores") or {}).items():
            h.pintar(r, c, col)
        if p.get("comentado"):
            h.comentar(r, "A", p.get("comentario_antes"))
        log(f"EXCEL  borro {p['hoja']} fila {r}")
    elif p["op"] == "proveedor_excel":
        h = libro.hoja("Proveedores")
        r = p["fila"]
        if str(h.formula(r, "B")) != p["valores"][1]:
            problemas.append(f"Proveedores fila {r} cambió después del lote; no la toco")
            return
        for c, v in zip(COLS_PROV, p["valores"]):
            if v:
                h.poner(r, c, "")
        log(f"EXCEL  borro Proveedores fila {r}")


def _revertir(libro, hechos) -> None:
    for p in reversed(hechos):
        try:
            _restaurar(libro, p, [], lambda s: None)
        except Exception:
            pass


def pasos_guardados(pasos: list[dict]) -> list[dict]:
    """Excel steps that were actually saved (followed by 'excel_guardado' before any 'excel_revertido')."""
    out, tanda = [], []
    for p in pasos:
        if p["op"] in ("fila_excel", "proveedor_excel"):
            tanda.append(p)
        elif p["op"] == "excel_guardado":
            out += tanda
            tanda = []
        elif p["op"] == "excel_revertido":
            tanda = []
    return out


def tanda_cortada(pasos: list[dict]) -> list[dict]:
    """Excel steps of the last run with neither 'excel_guardado' nor 'excel_revertido' after them: the app
    closed while writing, maybe right after Excel saved the workbook."""
    tanda: list[dict] = []
    for p in pasos:
        if p["op"] in ("fila_excel", "proveedor_excel"):
            tanda.append(p)
        elif p["op"] in ("excel_guardado", "excel_revertido"):
            tanda = []
    return tanda


def _sigue_como_la_dejo(libro, p: dict) -> bool:
    if p["op"] == "fila_excel":
        h = libro.hoja(p["hoja"])
        return str(h.formula(p["fila"], "A")) == str(p["escrito"]["A"]) and (
            "J" not in p["escrito"] or str(h.formula(p["fila"], "J")) == p["escrito"]["J"])
    return str(libro.hoja("Proveedores").formula(p["fila"], "B")) == p["valores"][1]


def reconocer_cortada(cfg, diario, log=print) -> bool | None:
    """Decides a cut run by looking at the cells, so finishing the batch never writes its rows twice.
    Excel saves all or nothing: if any row is still as the app wrote it, the run was saved. Returns True
    (saved), False (not saved) or None (nothing cut). Raises LibroConCambios if the workbook has unsaved
    changes (they may be the cut run itself, still in memory: the user saves or discards them first)."""
    tanda = tanda_cortada(diario.pasos)
    if not tanda:
        return None
    libro = abrir(cfg, Path(tanda[0]["libro"]))
    try:
        if not libro.saved:
            raise LibroConCambios("la planilla está abierta con cambios sin guardar: guardala o cerrala sin "
                                  "guardar y volvé a tocar Terminar este lote")
        guardada = any(_sigue_como_la_dejo(libro, p) for p in tanda)
    finally:
        libro.cerrar()
    diario.anotar({"op": "excel_guardado" if guardada else "excel_revertido", "recuperado": True})
    log(f"EXCEL la escritura cortada {'se había guardado' if guardada else 'no se había guardado'}")
    return guardada


def deshacer_excel(cfg, pasos: list[dict], log=print) -> list[str]:
    """Undo the saved Excel steps of a batch (newest first). Only rows still as the app left them."""
    pasos = pasos_guardados(pasos)
    if not pasos:
        return []
    problemas: list[str] = []
    libro = abrir(cfg, Path(pasos[0]["libro"]))
    try:
        if not libro.saved:
            raise LibroConCambios("la planilla está abierta con cambios sin guardar: guardala y volvé a deshacer")
        for p in reversed(pasos):
            _restaurar(libro, p, problemas, log)
        libro.guardar()
    finally:
        libro.cerrar()
    return problemas


# ======================================================================= clipboard

def html_portapapeles(lote) -> tuple[str, str]:
    """(HTML table, plain text) with the rows A:J, for pasting by hand if Excel cannot be opened."""
    import html
    css = "font-family:Arial;font-size:10pt"
    filas_html, plano = [], []
    for f in lote.filas:
        if f.excluida:
            continue
        celdas = [f.proveedor, f.tarea, "", f.empresa,
                  f.fecha.strftime("%d/%m/%Y") if f.fecha else "", f.vto.strftime("%d/%m/%Y") if f.vto else "",
                  f"{f.monto:.2f}".replace(".", ",") if f.monto is not None else "",
                  f"{f.iva:.2f}".replace(".", ",") if f.iva is not None else "", ""]
        tds = "".join(f'<td style="{css}{";mso-number-format:\'@\'" if i == 1 else ""}">{html.escape(str(v))}</td>'
                      for i, v in enumerate(celdas))
        link = f'<td style="{css}"><a href="{html.escape(f.url)}">{html.escape(f.texto_link)}</a></td>'
        filas_html.append(f"<tr>{tds}{link}</tr>")
        plano.append("\t".join(celdas + [f.url]))
    return "<table>" + "".join(filas_html) + "</table>", "\r\n".join(plano)


def copiar_portapapeles(lote) -> None:
    import win32clipboard
    tabla, plano = html_portapapeles(lote)
    pre = "<html><head><meta charset=\"utf-8\"></head><body><!--StartFragment-->"
    post = "<!--EndFragment--></body></html>"
    hdr = "Version:0.9\r\nStartHTML:{0:010d}\r\nEndHTML:{1:010d}\r\nStartFragment:{2:010d}\r\nEndFragment:{3:010d}\r\n"
    largo = len(hdr.format(0, 0, 0, 0).encode())
    sh = largo
    sf = sh + len(pre.encode())
    ef = sf + len(tabla.encode())
    eh = ef + len(post.encode())
    cf = (hdr.format(sh, eh, sf, ef) + pre + tabla + post).encode("utf-8")
    fmt = win32clipboard.RegisterClipboardFormat("HTML Format")
    win32clipboard.OpenClipboard()
    try:
        win32clipboard.EmptyClipboard()
        win32clipboard.SetClipboardData(fmt, cf)
        win32clipboard.SetClipboardText(plano, win32clipboard.CF_UNICODETEXT)
    finally:
        win32clipboard.CloseClipboard()
