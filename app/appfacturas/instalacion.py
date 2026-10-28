"""Install helpers: 'verificar' (what this PC has) and 'empaquetar' (zip to install on another PC)."""
from __future__ import annotations

import os
import sys
import zipfile
from datetime import datetime
from pathlib import Path

from . import NOMBRE, __version__, config

# what goes into the package for another PC (never tests: they have real invoices and bank details)
INCLUIR = ["app/appfacturas", "requirements.txt", "requirements-ui.txt", "requirements-nodeps.txt",
           "Instalar o actualizar RUA EZ-Fac.bat", "Abrir RUA EZ-Fac (con consola).bat",
           "Guia de uso RUA EZ-Fac.pdf", "Documentacion tecnica RUA EZ-Fac.md", "wheels"]


def _ok(b: bool) -> str:
    return "OK   " if b else "FALTA"


def verificar(salida: Path | None = None) -> bool:
    """Checks the PC can run the app. Returns False only for things that stop it from working."""
    ok, r = chequeo()
    texto = "\n".join(r)
    print(texto)
    if salida:
        salida.parent.mkdir(parents=True, exist_ok=True)
        salida.write_text(texto, encoding="utf-8")
    return ok


def chequeo() -> tuple[bool, list[str]]:
    r = [f"{NOMBRE} {__version__} - verificación {datetime.now():%d/%m/%Y %H:%M}", ""]
    grave = False
    r.append(f"OK    Python {sys.version.split()[0]}")
    if en_onedrive():
        grave = True
        r.append(f"FALTA la app está dentro de OneDrive ({config.carpeta_app()}): movela a Documentos y reinstalá")
    for mod, nombre in (("fitz", "PyMuPDF (lee los PDF)"), ("nicegui", "NiceGUI (pantallas)"),
                        ("openpyxl", "openpyxl (lee la planilla)"), ("win32com.client", "pywin32 (escribe en Excel)")):
        try:
            __import__(mod)
            r.append(f"OK    {nombre}")
        except Exception as ex:
            grave = grave or mod != "win32com.client" or sys.platform == "win32"
            r.append(f"FALTA {nombre}: {ex}")

    cfg = config.cargar()
    r.append("")
    if not cfg.base:
        grave = True
        r.append('FALTA carpeta "Facturas pendientes" en el OneDrive de RUA.')
        r.append('      Tiene que estar sincronizada en esta PC (OneDrive - Rua Asistencia\\Facturas pendientes),')
        r.append('      o poné la ruta en config.toml > [rutas] facturas_pendientes.')
    else:
        r.append(f"OK    Facturas pendientes: {cfg.base}")
        for sub in ("_Entrada\\Vehicular", "_Entrada\\Hogar", "_Entrada\\_Duplicados", "_Entrada\\_Procesados", "Vehicular", "Hogar"):
            p = cfg.base.joinpath(*sub.split("\\"))
            r.append(f"{_ok(p.is_dir())} {sub}" + ("" if p.is_dir() else "  (Setup la crea)"))
        r.append(f"{_ok(cfg.csv_proveedores.is_file())} _Script\\proveedores.csv"
                 + ("" if cfg.csv_proveedores.is_file() else "  (Setup lo crea vacío)"))
    if cfg.planilla and cfg.planilla.is_file():
        r.append(f"OK    Planilla: {cfg.planilla}")
        try:
            from openpyxl import load_workbook
            wb = load_workbook(cfg.planilla, read_only=True)
            for h in ("Vehicular", "Hogar", "Proveedores"):
                r.append(f"{_ok(h in wb.sheetnames)} hoja {h}")
            wb.close()
        except Exception as ex:
            r.append(f"AVISO no pude abrir la planilla para revisarla: {ex}")
    else:
        grave = True
        r.append(f"FALTA planilla {cfg.planilla or 'DEUDA A PAGAR ARG.xlsx'} (en la raíz del OneDrive de RUA; Setup crea una vacía)")
    r.append(f"{_ok(bool(cfg.url_base))} Links de SharePoint: {cfg.url_base or 'no pude armarlos (poné url_base en config.toml)'}")

    if sys.platform == "win32":
        try:
            import winreg
            winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "Excel.Application"))
            r.append("OK    Excel instalado")
        except OSError:
            grave = True
            r.append("FALTA Excel de escritorio (hace falta para escribir la planilla)")
        from .ui import _navegador_app
        nav = _navegador_app()
        r.append(f"{_ok(bool(nav))} Ventana propia: {nav or 'no encontré Edge ni Chrome; se abre en el navegador'}")
    r.append(f"OK    Modo real: {'HABILITADO' if cfg.modo_real else 'apagado (solo simula sobre lo real)'}")
    r.append("")
    faltan = any(l.startswith("FALTA") for l in r)
    r.append("Hay cosas que impiden usar la app (FALTA). Tocá Setup en Configuración o resolvelas a mano." if grave
             else "Casi listo: faltan algunas cosas (Setup las crea)." if faltan else "Todo listo.")
    return not grave, r


def empaquetar(destino: Path | None = None) -> Path:
    """Zip with what another PC needs. No venv, data, test folders, invoices or config."""
    raiz = config.carpeta_app()
    destino = destino or raiz / f"{NOMBRE} {__version__} para instalar.zip"
    tope = NOMBRE
    n = 0
    with zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as z:
        for item in INCLUIR:
            p = raiz / item
            if p.is_file():
                z.write(p, f"{tope}/{item}")
                n += 1
            elif p.is_dir():
                for f in sorted(p.rglob("*")):
                    if f.is_file() and "__pycache__" not in f.parts and f.suffix != ".pyc":
                        z.write(f, f"{tope}/{f.relative_to(raiz).as_posix()}")
                        n += 1
    print(f"Listo: {destino} ({n} archivos, {destino.stat().st_size // 1024} KB)")
    print("Pasale ese zip a tu compañero. Tiene que descomprimirlo en Documentos (no en OneDrive)")
    print('y hacer doble clic en "Instalar o actualizar RUA EZ-Fac.bat".')
    return destino


def en_onedrive() -> bool:
    return any(p.lower().startswith("onedrive") for p in config.carpeta_app().parts) or \
        bool(os.environ.get("OneDrive")) and str(config.carpeta_app()).lower().startswith(os.environ["OneDrive"].lower())
