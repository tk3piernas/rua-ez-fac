"""Read-only view of DEUDA A PAGAR ARG.xlsx with openpyxl (no Excel needed).

Writing goes through Excel COM (stage 4); reading for duplicates does not need it."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from urllib.parse import unquote

from .config import AREAS
from .maestro import _texto
from .util import solo_digitos

PRIMERA, ULTIMA = 12, 210
_HL_RE = re.compile(r'HYPERLINK\(\s*"([^"]*)"\s*[,;]\s*"([^"]*)"', re.I)


@dataclass
class FilaPlanilla:
    hoja: str
    fila: int | str            # sheet row, or "nueva" for rows of this batch
    proveedor: str
    tarea: str = ""
    fecha: date | None = None
    monto: Decimal | None = None
    j_texto: str = ""
    j_link: str = ""
    h_pisado: Decimal | None = None     # H written as a value instead of the IVA formula


@dataclass
class Planilla:
    ruta: Path | None
    filas: list[FilaPlanilla] = field(default_factory=list)
    proveedores: list[dict] = field(default_factory=list)
    alicuota: dict = field(default_factory=dict)      # sheet -> B4
    error: str = ""
    faltan: list = field(default_factory=list)         # area sheets the workbook doesn't have (Setup adds them)

    @property
    def leida(self) -> bool:
        return not self.error


def _fecha(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def _dec(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return Decimal(str(round(v, 2)))
    return None


def leer(ruta: Path | None, hojas=AREAS) -> Planilla:
    pl = Planilla(ruta)
    if not ruta or not ruta.exists():
        pl.error = f"no encontré la planilla ({ruta})"
        return pl
    try:
        import openpyxl
        wb = openpyxl.load_workbook(ruta, data_only=False, read_only=False)
    except Exception as e:
        pl.error = f"no pude abrir la planilla: {e}"
        return pl
    try:
        for h in hojas:
            if h not in wb.sheetnames:
                pl.faltan.append(h)
                continue
            ws = wb[h]
            b4 = ws["B4"].value
            pl.alicuota[h] = Decimal(str(b4)) if isinstance(b4, (int, float)) else Decimal("0.21")
            for r in range(PRIMERA, ULTIMA + 1):
                a = ws.cell(r, 1).value
                if a is None or not str(a).strip():
                    continue
                j = ws.cell(r, 10)
                texto, link = "", ""
                if isinstance(j.value, str) and j.value.startswith("="):
                    m = _HL_RE.search(j.value)
                    if m:
                        link, texto = m.group(1), m.group(2)
                else:
                    texto = "" if j.value is None else str(j.value)
                if j.hyperlink is not None and j.hyperlink.target:
                    link = link or j.hyperlink.target
                pl.filas.append(FilaPlanilla(h, r, str(a).strip(), str(ws.cell(r, 2).value or "").strip(),
                                             _fecha(ws.cell(r, 5).value), _dec(ws.cell(r, 7).value), texto, link,
                                             _dec(ws.cell(r, 8).value)))
        if "Proveedores" in wb.sheetnames:
            ws = wb["Proveedores"]
            for r in range(5, 301):
                vals = [ws.cell(r, c).value for c in range(1, 16)]
                nombre = str(vals[1] or "").strip()
                # typed by hand as a number: 20123456789.0 -> "20123456789" (B4)
                cuit = solo_digitos(_texto(vals[2]))
                if not nombre and not cuit:
                    continue
                pl.proveedores.append({"fila": r, "area": str(vals[0] or ""), "nombre": nombre, "cuit": cuit,
                                       "cbu": _cbu(vals[11])})
    finally:
        wb.close()
    return pl


def _cbu(v) -> str:
    """A CBU typed as a number lost its last digits (Excel keeps 15): better unknown than wrong."""
    if isinstance(v, float):
        return ""
    return solo_digitos(_texto(v))


def num_de(s: str) -> str | None:
    """'2-189' from 'Factura 0002-00000189' or a link ending in an AFIP file name (PS1 Num-De)."""
    if not s:
        return None
    s = unquote(s)
    hoja = re.split(r"[\\/]", s)[-1]
    m = re.match(r"^\d{11}_\d{3}_(\d{4,5})_(\d{8})", hoja)
    if m:
        return f"{int(m.group(1))}-{int(m.group(2))}"
    m = re.search(r"(?<!\d)(\d{1,5})-(\d{4,8})(?!\d)", s)
    if m:
        return f"{int(m.group(1))}-{int(m.group(2))}"
    return None


def tareas_de(s: str) -> list[str]:
    return re.findall(r"(?<!\d)\d{6}(?!\d)", s or "")
