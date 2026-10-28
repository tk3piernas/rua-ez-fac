"""Setup (create the local and OneDrive folders + an empty workbook if missing) and workbook reset."""
from __future__ import annotations

import os
import shutil
from datetime import datetime
from pathlib import Path

from . import config
from .excel import PRIMERA, ULTIMA, ExcelError, abrir, respaldar
from .maestro import COLS

AREAS = config.AREAS
SUBCARPETAS = ([f"_Entrada/{a}" for a in AREAS] + ["_Entrada/_Duplicados", "_Entrada/_Procesados"]
               + list(AREAS) + ["_Script"])
CSV_ENCABEZADO = "CUIT;Area;Proveedor;Carpeta\r\n"
NOMBRE_PLANILLA = "DEUDA A PAGAR ARG.xlsx"
COL_LISTA = "ABCDEFGH"                                     # Listas column with the companies of each area
H_R1C1 = '=IF(RC[-1]="","",ROUND(RC[-1]*R4C2,2))'          # H = G * B4 (IVA)
I_R1C1 = '=IF(RC[-2]="","",RC[-2]+RC[-1])'                 # I = G + H
COLS_PROVEEDORES = COLS
CARPETA_RESPALDOS = "Respaldos"                           # next to the workbook, in the OneDrive


# ------------------------------------------------------------------ where things go

def raiz_onedrive() -> Path | None:
    """The synced OneDrive for Business folder of RUA ('OneDrive - Rua Asistencia')."""
    cands = []
    if os.environ.get("OneDriveCommercial"):
        cands.append(Path(os.environ["OneDriveCommercial"]))
    home = Path(os.environ.get("USERPROFILE") or Path.home())
    if home.exists():
        todas = sorted(p for p in home.glob("OneDrive - *") if p.is_dir())
        cands += [p for p in todas if "rua" in p.name.lower()] + todas
    return next((c for c in cands if c.is_dir()), None)


def plan_setup(cfg) -> dict:
    """What setup would create. Never overwrites anything."""
    base = cfg.base
    raiz = base.parent if base else raiz_onedrive()
    if not base and raiz:
        base = raiz / "Facturas pendientes"
    crear_dirs, crear_archivos = [], []
    app = config.carpeta_app()
    for d in ("datos", "datos/lotes", "datos/backups", "datos/entrada_ui"):
        if not (app / d).is_dir():
            crear_dirs.append(app / d)
    if base:
        if not base.is_dir():
            crear_dirs.append(base)
        for s in SUBCARPETAS:
            if not (base / s).is_dir():
                crear_dirs.append(base / s)
        if not (base / "_Script" / "proveedores.csv").is_file():
            crear_archivos.append(base / "_Script" / "proveedores.csv")
    planilla = cfg.planilla or (raiz / NOMBRE_PLANILLA if raiz else None)
    # a missing workbook is asked apart: OneDrive may not have downloaded it yet (A9)
    nueva = planilla if planilla and not planilla.is_file() else None
    hojas = hojas_faltantes(planilla) if planilla and planilla.is_file() else []
    return {"base": base, "raiz": raiz, "planilla": planilla, "dirs": crear_dirs, "archivos": crear_archivos,
            "hojas": hojas, "sin_onedrive": base is None, "planilla_nueva": nueva,
            "onedrive_sin_bajar": onedrive_sin_bajar(planilla.parent if planilla else raiz) if nueva else ""}


def onedrive_sin_bajar(carpeta: Path | None) -> str:
    """Why the OneDrive folder may still be syncing ("" if it looks complete): empty, or with files that are
    only in the cloud. Then a missing workbook may just not be here yet."""
    if not carpeta or not carpeta.is_dir():
        return "la carpeta del OneDrive todavía no existe"
    try:
        hijos = list(carpeta.iterdir())
    except OSError:
        return "no puedo ver la carpeta del OneDrive"
    if not hijos:
        return "la carpeta del OneDrive está vacía"
    SOLO_NUBE = 0x00400000 | 0x00040000 | 0x00001000   # RECALL_ON_DATA_ACCESS, RECALL_ON_OPEN, OFFLINE
    for h in hijos:
        try:
            if getattr(h.stat(), "st_file_attributes", 0) & SOLO_NUBE:
                return f"hay archivos que OneDrive todavía no bajó (por ejemplo {h.name})"
        except OSError:
            continue
    return ""


def ejecutar_setup(cfg, plan: dict | None = None, crear_planilla: bool = False) -> list[str]:
    """Creates what the plan says. The empty workbook only with `crear_planilla` (its own confirmation)."""
    plan = plan or plan_setup(cfg)
    hecho = []
    if crear_planilla and plan.get("planilla_nueva") and not plan["planilla_nueva"].exists():
        crear_planilla_vacia(plan["planilla_nueva"])
        hecho.append(f"Planilla nueva vacía: {plan['planilla_nueva']}")
    for d in plan["dirs"]:
        d.mkdir(parents=True, exist_ok=True)
        hecho.append(f"Carpeta creada: {d}")
    for f in plan["archivos"]:
        if f.exists():
            continue
        f.parent.mkdir(parents=True, exist_ok=True)
        if f.name == "proveedores.csv":
            f.write_text(CSV_ENCABEZADO, encoding="utf-8-sig", newline="")
        elif f.suffix.lower() == ".xlsx":
            crear_planilla_vacia(f)
        hecho.append(f"Archivo creado: {f}")
    if plan.get("hojas"):
        cfg_p = cfg if cfg.planilla else config.Config(base=plan["base"], planilla=plan["planilla"],
                                                       excel_motor=cfg.excel_motor)
        res = agregar_hojas(cfg_p, config.carpeta_app() / "datos" / "backups")
        hecho += [f"Hoja agregada a la planilla: {h}" for h in res["hojas"]]
    return hecho


# ------------------------------------------------------------------ area sheets missing in the workbook

def hojas_faltantes(planilla: Path) -> list[str]:
    """Area sheets the workbook doesn't have yet (e.g. a workbook from before Persona / ServiciosYa)."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(planilla, read_only=True)
        nombres = wb.sheetnames
        wb.close()
    except Exception:
        return []
    return [a for a in AREAS if a not in nombres]


def agregar_hojas(cfg, backups: Path, log=print) -> dict:
    """Adds the missing area sheets as empty copies of the last existing area sheet (same layout, formulas
    and summary). Before, a backup copy of the workbook (to undo: put the copy back)."""
    libro = abrir(cfg)
    try:
        if not libro.saved:
            raise ExcelError("la planilla está abierta con cambios sin guardar: guardala o cerrala y volvé a probar")
        nombres = libro.nombres()
        faltan = [a for a in AREAS if a not in nombres]
        if not faltan:
            return {"hojas": [], "respaldos": []}
        modelo = next((a for a in reversed(AREAS) if a in nombres), None)
        if not modelo:
            raise ExcelError("la planilla no tiene ninguna hoja de área para copiar")
        copias = respaldo(Path(cfg.planilla), backups, "antes de agregar hojas")
        for area in faltan:
            h = libro.copiar_hoja(modelo, area)
            titulo = h.valor(1, "A")
            if isinstance(titulo, str) and modelo in titulo:
                h.poner(1, "A", titulo.replace(modelo, area))
            ref = next((r for r in range(PRIMERA, ULTIMA + 1) if str(h.formula(r, "H")).startswith("=")), None)
            h.vaciar_rango(f"A{PRIMERA}:J{ULTIMA}", sin_columna="I", formula_h_desde=ref)
            col = COL_LISTA[AREAS.index(area)]
            if "Listas" in nombres:
                hl = libro.hoja("Listas")
                if hl.valor(1, col) in ("", None):
                    hl.poner(1, col, f"Empresas {area}")
                h.lista(f"D{PRIMERA}:D{ULTIMA}", f"=Listas!${col}$2:${col}$50")
            log(f"Hoja {area} agregada (copia vacía de {modelo})")
            modelo = area
        libro.guardar()
        return {"hojas": faltan, "respaldos": [str(c) for c in copias]}
    finally:
        libro.cerrar()


# ------------------------------------------------------------------ new empty workbook

def crear_planilla_vacia(ruta: Path) -> Path:
    """DEUDA A PAGAR ARG layout: Vehicular/Hogar (headers row 11, data 12-210, H/I formulas, summary H4:I9),
    Proveedores (headers row 4, A-O) and Listas."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    navy, naranja = "14147A", "FF542C"
    titulo = Font(bold=True, size=14, color=navy)
    enc = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor=navy)
    fino = Side(style="thin", color="D0D4E0")
    borde = Border(bottom=fino)
    wb = Workbook()
    wb.remove(wb.active)
    for i, area in enumerate(AREAS):
        ws = wb.create_sheet(area)
        ws["A1"] = f"Deuda a pagar · {area}"
        ws["A1"].font = titulo
        ws["A4"], ws["B4"] = "Alícuota IVA", 0.21
        ws["B4"].number_format = "0%"
        ws["A5"], ws["B5"] = "Hoy", "=TODAY()"
        ws["B5"].number_format = "dd/mm/yyyy"
        rng = lambda c: f"{c}{PRIMERA}:{c}{ULTIMA}"
        resumen = [("Total pendiente", f"=SUM({rng('G')})"), ("Total IVA", f"=SUM({rng('H')})"),
                   ("Total + IVA", f"=SUM({rng('I')})"), ("# de facturas", f"=COUNTA({rng('A')})"),
                   ("Vencido", f'=SUMIFS({rng("I")},{rng("F")},"<"&$B$5)'),
                   ("Vencen pronto (7 días)", f'=SUMIFS({rng("I")},{rng("F")},">="&$B$5,{rng("F")},"<="&$B$5+7)')]
        for j, (t, f) in enumerate(resumen):
            ws.cell(4 + j, 8, t).font = Font(bold=True, color=navy)
            c = ws.cell(4 + j, 9, f)
            c.number_format = "0" if t.startswith("#") else '"$" #,##0.00'
        heads = ["Proveedor", "N° de tarea", "Servicio realizado", "Empresa", "Fecha realizada",
                 "Fecha de vencimiento", "Monto total", "IVA", "Monto total + IVA", "Link a la factura"]
        for j, h in enumerate(heads, start=1):
            c = ws.cell(11, j, h)
            c.font, c.fill = enc, relleno
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[11].height = 30
        for r in range(PRIMERA, ULTIMA + 1):
            ws.cell(r, 2).number_format = "@"
            for c in (5, 6):
                ws.cell(r, c).number_format = "dd/mm/yyyy"
            for c in (7, 8, 9):
                ws.cell(r, c).number_format = '"$" #,##0.00'
            ws.cell(r, 8, f'=IF(G{r}="","",ROUND(G{r}*$B$4,2))')
            ws.cell(r, 9, f'=IF(G{r}="","",G{r}+H{r})')
            for c in range(1, 11):
                ws.cell(r, c).border = borde
        for col, w in zip("ABCDEFGHIJ", (34, 13, 34, 18, 14, 14, 15, 13, 17, 30)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A12"
        ws.conditional_formatting.add(f"A{PRIMERA}:J{ULTIMA}", FormulaRule(
            formula=[f'AND($F{PRIMERA}<>"",$F{PRIMERA}<$B$5,$A{PRIMERA}<>"")'], fill=PatternFill("solid", fgColor="FDE2E1")))
        ws.conditional_formatting.add(f"A{PRIMERA}:J{ULTIMA}", FormulaRule(
            formula=[f'AND($F{PRIMERA}<>"",$F{PRIMERA}>=$B$5,$F{PRIMERA}<=$B$5+7,$A{PRIMERA}<>"")'],
            fill=PatternFill("solid", fgColor="FFF6D5")))
        dv = DataValidation(type="list", formula1=f"=Listas!${COL_LISTA[i]}$2:${COL_LISTA[i]}$50", allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"D{PRIMERA}:D{ULTIMA}")
    wp = wb.create_sheet("Proveedores")
    wp["A1"] = "Proveedores"
    wp["A1"].font = titulo
    for j, h in enumerate(COLS_PROVEEDORES, start=1):
        c = wp.cell(4, j, h)
        c.font, c.fill = enc, relleno
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        wp.column_dimensions[c.column_letter].width = 16 if j not in (2, 5, 15) else 30
    for r in range(5, 301):
        for c in (3, 7, 11, 12):
            wp.cell(r, c).number_format = "@"
    wp.freeze_panes = "A5"
    wl = wb.create_sheet("Listas")
    for i, area in enumerate(AREAS):
        c = wl[f"{COL_LISTA[i]}1"]
        c.value, c.font = f"Empresas {area}", Font(bold=True, color=naranja)
        wl.column_dimensions[COL_LISTA[i]].width = 28
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return ruta


# ------------------------------------------------------------------ reset

def respaldo_reset(planilla: Path, backups: Path) -> list[Path]:
    return respaldo(planilla, backups, "antes del reset")


def respaldo(planilla: Path, backups: Path, motivo: str) -> list[Path]:
    """Two copies before a big change: in 'Respaldos' next to the workbook (OneDrive keeps it) and in
    datos\\backups (only the newest 20 stay there)."""
    marca = f"{datetime.now():%Y-%m-%d %H%M}"
    carpeta = planilla.parent / CARPETA_RESPALDOS
    carpeta.mkdir(exist_ok=True)
    junto = carpeta / f"{planilla.stem} - respaldo {motivo} {marca}{planilla.suffix}"
    shutil.copy2(planilla, junto)
    local = respaldar(planilla, backups, motivo)
    return [junto, local]


def reset_planilla(cfg, backups: Path, vaciar_proveedores: bool = False, log=print, solo_local: bool = False) -> dict:
    """Empties the workbook IN PLACE (same file, so the OneDrive sharing links keep working):
    A-G and J of rows 12-210 (values, fills, comments, links) and H back to its formula. I is never touched.
    Optionally empties Proveedores A5:O300. Returns what was done (with the backups).
    `solo_local`: only the copy in datos\backups (closing a batch round already left an exported copy in OneDrive)."""
    if not cfg.planilla or not Path(cfg.planilla).is_file():
        raise ExcelError(f"no encuentro la planilla ({cfg.planilla})")
    libro = abrir(cfg)
    try:
        if not libro.saved:
            raise ExcelError("la planilla está abierta con cambios sin guardar: guardala o cerrala y volvé a probar")
        if solo_local:
            local = respaldar(Path(cfg.planilla), backups, "antes del reset")
            copias = [local, local]
        else:
            copias = respaldo_reset(Path(cfg.planilla), backups)
        filas, provs = vaciar_libro(libro, vaciar_proveedores, log)
        libro.guardar()
    finally:
        libro.cerrar()
    _anotar_reset(Path(cfg.planilla), backups, copias)
    return {"filas": filas, "proveedores": provs, "respaldos": [str(c) for c in dict.fromkeys(copias)]}


def contar_filas(libro) -> tuple[dict, int]:
    """Rows with data in each area sheet and in Proveedores (what a reset would empty)."""
    nombres = libro.nombres()
    filas = {}
    for area in (a for a in AREAS if a in nombres):
        h = libro.hoja(area)
        filas[area] = sum(1 for r in range(PRIMERA, ULTIMA + 1) if h.formula(r, "A") not in ("", None))
    provs = 0
    if "Proveedores" in nombres:
        hp = libro.hoja("Proveedores")
        provs = sum(1 for r in range(5, 301) if hp.formula(r, "A") not in ("", None) or hp.formula(r, "B") not in ("", None))
    return filas, provs


def vaciar_libro(libro, vaciar_proveedores: bool, log=print) -> tuple[dict, int]:
    """The reset itself on an open workbook (not saved here): area sheets rows 12-210 without column I,
    and optionally Proveedores A5:O300. Returns the rows emptied."""
    filas, provs = contar_filas(libro)
    for area in filas:
        h = libro.hoja(area)
        ref = next((r for r in range(PRIMERA, ULTIMA + 1) if str(h.formula(r, "H")).startswith("=")), None)
        h.vaciar_rango(f"A{PRIMERA}:J{ULTIMA}", sin_columna="I", formula_h_desde=ref)
        log(f"{area}: {filas[area]} filas vaciadas")
    if vaciar_proveedores:
        libro.hoja("Proveedores").limpiar("A5:O300")
        log(f"Proveedores: {provs} filas vaciadas")
    else:
        provs = 0
    return filas, provs


def _archivo_reset(backups: Path) -> Path:
    return backups.parent / "ultimo_reset.json"


def _anotar_reset(planilla: Path, backups: Path, copias: list[Path]) -> None:
    import json
    from .util import sha256
    _archivo_reset(backups).write_text(json.dumps(
        {"fecha": datetime.now().isoformat(timespec="seconds"), "planilla": str(planilla),
         "respaldo": str(copias[0]), "respaldo_local": str(copias[1]),
         "sha256": sha256(planilla), "deshecho": False}, ensure_ascii=False, indent=1), encoding="utf-8")


def reset_para_volver(cfg, backups: Path) -> dict | None:
    """The last reset of this workbook, if it can still be undone: nobody changed the workbook since
    (same SHA256) and its backup is still there."""
    import json
    from .util import sha256
    try:
        r = json.loads(_archivo_reset(backups).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    pl = Path(cfg.planilla) if cfg.planilla else None
    if r.get("deshecho") or not pl or not pl.is_file() or str(pl).lower() != r["planilla"].lower():
        return None
    copia = next((Path(c) for c in (r["respaldo"], r.get("respaldo_local", "")) if c and Path(c).is_file()), None)
    if not copia or sha256(pl) != r["sha256"]:
        return None
    return {**r, "copia": str(copia)}


def volver_del_reset(cfg, backups: Path, log=print) -> Path:
    """Puts back the workbook as it was before the last reset. Same file (written in place, so the OneDrive
    sharing links keep working). Only if nobody changed it since the reset, and with the workbook closed."""
    import json
    r = reset_para_volver(cfg, backups)
    if not r:
        raise ExcelError("no hay un reset para deshacer (o la planilla cambió después del reset)")
    pl = Path(cfg.planilla)
    if _abierta_en_excel(cfg, pl):
        raise ExcelError("la planilla está abierta en Excel: cerrala y volvé a probar")
    shutil.copyfile(r["copia"], pl)
    r["deshecho"] = True
    r.pop("copia")
    _archivo_reset(backups).write_text(json.dumps(r, ensure_ascii=False, indent=1), encoding="utf-8")
    log(f"PLANILLA vuelve a como estaba antes del reset del {r['fecha']}")
    return pl


def _abierta_en_excel(cfg, pl: Path) -> bool:
    if getattr(cfg, "excel_motor", "com") != "com":
        return False
    try:
        from .excel import LibroCOM, url_onedrive_de
        return LibroCOM._buscar_abierto(pl, url_onedrive_de(cfg, pl)) is not None
    except Exception:
        return False


# ------------------------------------------------------------------ export

CARPETA_EXPORTADAS = "Planillas exportadas"


def nombre_exportacion(nombre: str) -> str:
    """A valid Windows file name ending in .xlsx (or "" if nothing is left)."""
    import re
    n = re.sub(r'[<>"|?*\x00-\x1f]', "", re.sub(r"[/\\:]", "-", nombre or "")).strip().rstrip(".")
    if n.lower().endswith(".xlsx"):
        n = n[:-5].rstrip()
    return f"{n}.xlsx" if n else ""


def exportar_planilla(cfg, nombre: str, log=print) -> Path:
    """Saves a copy of the workbook (as Excel has it now) in 'Planillas exportadas', next to it.
    Never overwrites: if the name exists it says so. Nothing else changes."""
    if not cfg.planilla or not Path(cfg.planilla).is_file():
        raise ExcelError(f"no encuentro la planilla ({cfg.planilla})")
    archivo = nombre_exportacion(nombre)
    if not archivo:
        raise ExcelError("poné un nombre para la planilla exportada")
    destino = Path(cfg.planilla).parent / CARPETA_EXPORTADAS / archivo
    if destino.exists():
        raise ExcelError(f"ya hay una planilla exportada con ese nombre ({destino.name}); elegí otro")
    destino.parent.mkdir(parents=True, exist_ok=True)
    libro = abrir(cfg)
    try:
        libro.guardar_copia(destino)
    finally:
        libro.cerrar()
    log(f"EXPORTADA {destino}")
    return destino


def borrar_exportacion(destino: Path, sha: str) -> str:
    """Undo of an export: deletes the copy if nobody changed it. Returns a problem or ""."""
    from .util import sha256
    if not destino.exists():
        return ""
    if sha256(destino) != sha:
        return f"{destino.name} cambió después de exportarla; no la borro"
    destino.unlink()
    return ""
