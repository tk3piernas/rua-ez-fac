"""Automatic test of stage 3 on the sandbox: snapshot, real run, snapshot, undo, compare."""
from __future__ import annotations

from pathlib import Path

from . import ejecutor, paridad, reporte
from .config import carpeta_app, cargar
from .lote import Armador, entradas_de_carpeta
from .util import ahora_id, sha256


def foto(raiz: Path) -> dict[str, str]:
    """relative path -> sha256 (files) or '<dir>' (folders)."""
    out = {}
    for p in sorted(raiz.rglob("*")):
        out[str(p.relative_to(raiz))] = "<dir>" if p.is_dir() else sha256(p)
    return out


def correr_lote(cfg, simular: bool, log=print):
    lid = ahora_id()
    trabajo = carpeta_app() / "datos" / "lotes" / lid
    ents, txts = entradas_de_carpeta(cfg)
    lote = Armador(cfg, simular=simular, trabajo=trabajo).procesar(ents, txts)
    reporte.guardar(lote, trabajo.parent, reporte.texto(lote, cfg))
    return lote, trabajo


def prueba_mover(salida: Path) -> bool:
    L = []
    w = lambda s="": (L.append(s), print(s))
    base = paridad.preparar()
    raiz = base.parent
    f0 = foto(raiz)
    cfg = cargar(base=base)
    lote, trabajo = correr_lote(cfg, simular=False)
    w(f"Lote {lote.id}: {len(lote.facturas)} facturas, {len(lote.txts)} .txt")
    res = ejecutor.Ejecutor(cfg, lote, trabajo, log=w, con_excel=False).ejecutar()
    f1 = foto(raiz)
    nuevos = sorted(k for k in f1 if k not in f0)
    idos = sorted(k for k in f0 if k not in f1)
    w("")
    w(f"Después del lote: {len(nuevos)} archivos o carpetas nuevos, {len(idos)} que ya no están donde estaban")
    for k in nuevos:
        w(f"  + {k}")
    for k in idos:
        w(f"  - {k}")
    quedan = sorted(k for k in f1 if k.startswith(str(Path("Facturas pendientes") / "_Entrada"))
                    and f1[k] != "<dir>" and "_Duplicados" not in k and "_Procesados" not in k)
    w("")
    w("Quedan en _Entrada (no se movieron a propósito):")
    for k in quedan:
        w(f"  · {k}")
    csv0 = f0.get(str(Path("Facturas pendientes") / "_Script" / "proveedores.csv"))
    csv1 = f1.get(str(Path("Facturas pendientes") / "_Script" / "proveedores.csv"))
    w(f"proveedores.csv {'cambió (proveedores nuevos)' if csv0 != csv1 else 'sin cambios'}")
    if res["errores"]:
        w("ERRORES:")
        for e in res["errores"]:
            w(f"  ! {e}")
    w(f"Filas para la planilla que quedan pendientes (se escriben en la etapa 4): {res['excel_pendiente']}")
    w("")
    w("Deshaciendo el lote ...")
    problemas = ejecutor.deshacer(trabajo, log=w)
    f2 = foto(raiz)
    distinto = sorted(set(f0) ^ set(f2) | {k for k in f0 if k in f2 and f0[k] != f2[k]})
    w("")
    ok = not problemas and not distinto and not res["errores"]
    if ok:
        w("RESULTADO: OK. Después de deshacer, la carpeta de prueba quedó idéntica a como estaba (comparada archivo por archivo).")
    else:
        w("RESULTADO: HAY DIFERENCIAS")
        for p in problemas:
            w(f"  ! {p}")
        for k in distinto:
            w(f"  ≠ {k}")
    salida.write_text("\n".join(L), encoding="utf-8")
    return ok


# ---------------------------------------------------------------------------- stage 4

def foto_libro(ruta: Path, hojas=("Vehicular", "Hogar", "Proveedores"), filas=300, cols=15) -> dict:
    """Cell contents (values/formulas), fill and comment of the workbook, read from the saved file."""
    import openpyxl
    wb = openpyxl.load_workbook(ruta)
    out = {}
    try:
        for h in hojas:
            if h not in wb.sheetnames:
                continue
            for row in wb[h].iter_rows(min_row=1, max_row=filas, max_col=cols):
                for c in row:
                    fill = c.fill.fgColor.rgb if c.fill is not None and c.fill.fill_type else None
                    com = c.comment.text if c.comment else None
                    if c.value is not None or fill or com:
                        out[(h, c.coordinate)] = (c.value, fill, com)
    finally:
        wb.close()
    return out


def _abrir_visible(libro: Path, w) -> bool:
    """Opens the sandbox workbook in Excel like a person would, and waits until it is open."""
    import os
    import time
    from .excel import LibroCOM
    os.startfile(str(libro))
    for _ in range(60):
        time.sleep(1)
        try:
            if LibroCOM._buscar_abierto(libro, ""):
                time.sleep(2)
                return True
        except Exception:
            pass
    w("  ! Excel no abrió la planilla de prueba en 60 segundos")
    return False


def _cerrar_visible(libro: Path) -> None:
    from .excel import LibroCOM
    wb = LibroCOM._buscar_abierto(libro, "")
    if wb is not None:
        xl = wb.Application
        wb.Close(False)
        if xl.Workbooks.Count == 0:
            xl.Quit()


def prueba_excel(salida: Path) -> bool:
    from . import excel
    L = []
    w = lambda s="": (L.append(s), print(s))
    todo_ok = True
    for modo in ("cerrada", "abierta"):
        w(f"===== Planilla {modo} =====")
        base = paridad.preparar()
        libro = base.parent / "DEUDA A PAGAR ARG.xlsx"
        if not libro.exists():
            w("  ! No hay copia de la planilla en la carpeta de prueba")
            return False
        x0 = foto_libro(libro)
        if modo == "abierta" and not _abrir_visible(libro, w):
            todo_ok = False
            continue
        cfg = cargar(base=base)
        lote, trabajo = correr_lote(cfg, simular=False)
        res = ejecutor.Ejecutor(cfg, lote, trabajo, log=lambda s: None).ejecutar()
        w(f"Archivos: {len(res['hechos'])} pasos, {len(res['errores'])} errores")
        w(f"Excel: {res['excel']} {res['excel_mensaje']}  ({res['excel_escrito']['filas']} filas, "
          f"{res['excel_escrito']['proveedores']} proveedores)")
        ok = res["excel"] == "escrito" and not res["errores"]
        if modo == "abierta":
            wb = excel.LibroCOM._buscar_abierto(libro, "")
            sigue = wb is not None
            w(f"La planilla sigue abierta: {'sí' if sigue else 'NO'}; guardada: {'sí' if sigue and wb.Saved else 'NO'}")
            ok = ok and sigue and bool(wb.Saved)
        # check every written row against the saved file
        x1 = foto_libro(libro)
        pasos = [p for p in ejecutor.Diario(trabajo).pasos if p["op"] in ("fila_excel", "proveedor_excel")]
        malos = []
        for p in pasos:
            if p["op"] == "fila_excel":
                h, r = p["hoja"], p["fila"]
                val = lambda c: (x1.get((h, f"{c}{r}")) or (None,))[0]
                if val("A") != p["escrito"]["A"]:
                    malos.append(f"{h} fila {r}: A = {val('A')!r}")
                if not str(val("I") or "").startswith("="):
                    malos.append(f"{h} fila {r}: la fórmula de I se perdió")
                if "H" not in p["escrito"] and not str(val("H") or "").startswith("="):
                    malos.append(f"{h} fila {r}: la fórmula de H se perdió")
                if "J" in p["escrito"] and not str(val("J") or "").upper().startswith("=HYPERLINK("):
                    malos.append(f"{h} fila {r}: J no tiene el HYPERLINK")
                if "B" in p["escrito"] and val("B") != p["escrito"]["B"]:
                    malos.append(f"{h} fila {r}: B = {val('B')!r} (tendría que ser texto {p['escrito']['B']!r})")
                w(f"  {h} fila {r}: {p['escrito']['A']} | B {p['escrito'].get('B', '')} | G {p['escrito'].get('G', '')} | "
                  f"H {'valor ' + str(p['escrito']['H']) if 'H' in p['escrito'] else 'fórmula'}")
            else:
                r = p["fila"]
                for c, v in zip("ABCDEFGHIJKLMNO", p["valores"]):
                    got = (x1.get(("Proveedores", f"{c}{r}")) or (None,))[0]
                    if v and str(got) != v:
                        malos.append(f"Proveedores fila {r} columna {c}: {got!r} en vez de {v!r}")
                w(f"  Proveedores fila {r}: {p['valores'][1]} | CUIT {p['valores'][2]} | CBU {p['valores'][11]}")
        for m in malos:
            w(f"  ! {m}")
        ok = ok and not malos
        problemas = ejecutor.deshacer(trabajo, log=lambda s: None, cfg=cfg)
        if modo == "abierta":
            _cerrar_visible(libro)
        x2 = foto_libro(libro)
        dif = sorted(set(x0) ^ set(x2) | {k for k in x0 if k in x2 and x0[k] != x2[k]})
        for p in problemas:
            w(f"  ! deshacer: {p}")
        for k in dif[:20]:
            w(f"  ≠ {k[0]}!{k[1]}: antes {x0.get(k)!r} / después {x2.get(k)!r}")
        ok = ok and not problemas and not dif
        w(f"Deshacer: {'la planilla quedó idéntica (celda por celda)' if not dif and not problemas else 'HAY DIFERENCIAS'}")
        w(f"RESULTADO planilla {modo}: {'OK' if ok else 'CON PROBLEMAS'}")
        w("")
        todo_ok = todo_ok and ok
    salida.write_text("\n".join(L), encoding="utf-8")
    return todo_ok
