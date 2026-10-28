"""G1: the audit (Revisar) finds broken links, repeated rows, rows without link, PDFs without row, old files."""
import os
import time

from openpyxl import load_workbook

from appfacturas import auditoria, ejecutor, excel
from test_ejecutor_medias import _armar
from test_lote_sintetico import factura, sandbox  # noqa: F401  (fixture)


def _por_titulo(hallazgos):
    return {h.titulo.split(" hace")[0].split(" (")[0]: h.items for h in hallazgos}


def test_revisar_encuentra_lo_que_hay_que_mirar(sandbox, tmp_path):
    factura(sandbox.entrada / "Vehicular" / "uno.pdf", 111)
    factura(sandbox.entrada / "Vehicular" / "dos.pdf", 112)
    lote, trabajo = _armar(sandbox, tmp_path)
    ejecutor.Ejecutor(sandbox, lote, trabajo, log=lambda s: None).ejecutar()
    assert all(not h.items for h in auditoria.revisar(sandbox) if not h.titulo.startswith("Filas con el IVA"))

    wb = load_workbook(sandbox.planilla)
    ws = wb["Vehicular"]
    r = excel.PRIMERA
    link = ws[f"J{r}"].value.split('"')[1]
    borrado = auditoria.ruta_de_link(sandbox, link)
    borrado.unlink()                                                              # broken link (row 12)
    carpeta = borrado.parent
    (carpeta / "Factura 0003-00000999.pdf").write_bytes(b"x")                     # PDF without a row
    for c in "ABCDEFGJ":                                                          # same row twice
        ws[f"{c}{r + 2}"] = ws[f"{c}{r + 1}"].value
    ws[f"A{r + 3}"], ws[f"G{r + 3}"] = "Prov Sin Ficha", 500                       # no link, no payment details
    wb.save(sandbox.planilla)
    viejo = sandbox.entrada / "Hogar" / "olvidado.pdf"
    viejo.write_bytes(b"x")
    hace_10_dias = time.time() - 10 * 86400
    os.utime(viejo, (hace_10_dias, hace_10_dias))

    h = _por_titulo(auditoria.revisar(sandbox))
    assert len(h["Links a archivos que no están"]) == 1
    assert len(h["Filas repetidas"]) == 1
    assert len(h["Filas sin link"]) == 1
    assert h["PDF guardados que no están en la planilla"] == [str(carpeta.relative_to(sandbox.base) / "Factura 0003-00000999.pdf")]
    assert h["Proveedores sin datos de pago"] == ["Prov Sin Ficha"]
    assert len(h["Archivos esperando"]) == 1
    texto = auditoria.informe(auditoria.revisar(sandbox))
    assert "Filas repetidas: 1" in texto


def test_revisar_sin_planilla(sandbox):
    sandbox.planilla.unlink()
    (h,) = auditoria.revisar(sandbox)
    assert h.items and "No pude leer" in h.titulo
