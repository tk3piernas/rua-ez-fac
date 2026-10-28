"""'Borrar todo' with made-up data (no real invoices: this file goes to the repo).
Decoys in the OneDrive root and in the app must never move, change or disappear."""
import os
from pathlib import Path

import pytest
from openpyxl import load_workbook

from appfacturas import borrado, preparar
from appfacturas.config import Config
from appfacturas.util import sha256

OD = "OneDrive - Rua Asistencia"
CSV = "CUIT;Area;Proveedor;Carpeta\r\n20111111112;Vehicular;Remolques Inventados;Remolques Inventados\r\n"

# never the app's: they must stay exactly as they are
SENUELOS_OD = ["Escritorio/notas.txt", "Imágenes/foto.jpg", "Archivos adjuntos/adjunto.pdf",
               "Microsoft Teams Chat Files/chat.pdf", "Presupuesto 2026.xlsx", "DEUDA A PAGAR ARG (copia).xlsx",
               "DEUDA A PAGAR ARG - respaldo.docx", "Facturas pendientes viejas/vieja.pdf",
               "Respaldos personales/mio.xlsx", "Respaldos/mio.docx", "Planillas exportadas/Leeme.txt"]
SENUELOS_APP = ["config.toml", ".venv/pyvenv.cfg", "app/appfacturas/x.py", "datos/recorrido.json", "datos/app.log",
                "Otra cosa/archivo.txt"]


def _escribir(p: Path, texto: str) -> Path:
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(texto.encode("utf-8"))
    return p


def _arbol(tmp_path, monkeypatch, raiz: Path):
    """App folder (APPFACTURAS_HOME) and an OneDrive root with the app's things and decoys."""
    app = tmp_path / "app"
    monkeypatch.setenv("APPFACTURAS_HOME", str(app))
    base = raiz / "Facturas pendientes"
    for rel in ("Vehicular/Remolques Inventados/Factura 0001-00000001.pdf",
                "Vehicular/Remolques Inventados/Factura 0001-00000002.pdf",
                "Hogar/Cerrajería Ficticia/Factura 0002-00000010.pdf", "Hogar/Cerrajería Ficticia/datos.txt",
                "_Entrada/Hogar/nueva.pdf", "_Entrada/_Duplicados/dup.pdf", "_Entrada/_Procesados/vieja.pdf",
                "_Revisar con Claude/lote1/pedido.json"):
        _escribir(base / rel, f"contenido de {rel}")
    (base / "_Entrada" / "Vehicular").mkdir(parents=True)
    _escribir(base / "_Script" / "proveedores.csv", CSV)
    _escribir(raiz / "Planillas exportadas" / "Cierre octubre.xlsx", "export")
    _escribir(raiz / "Respaldos" / "DEUDA A PAGAR ARG - respaldo antes del reset 2026-10-01 1010.xlsx", "r1")
    _escribir(raiz / "PROVEEDORES RUA.xlsx", "memoria")
    _escribir(raiz / "DEUDA A PAGAR ARG - respaldo 2026-09-01.xlsx", "r0")
    pl = raiz / "DEUDA A PAGAR ARG.xlsx"
    preparar.crear_planilla_vacia(pl)
    wb = load_workbook(pl)
    wb["Vehicular"]["A12"], wb["Vehicular"]["G12"] = "Remolques Inventados", 1000
    wb["Hogar"]["A12"], wb["Hogar"]["A13"] = "Cerrajería Ficticia", "Otra"
    wb["Proveedores"]["A5"], wb["Proveedores"]["B5"] = "Vehicular", "Remolques Inventados"
    wb.save(pl)
    for s in SENUELOS_OD:
        _escribir(raiz / s, f"señuelo {s}")
    for s in SENUELOS_APP:
        _escribir(app / s, f"señuelo {s}")
    for rel in ("datos/lotes/L1/diario.json", "datos/tandas.json", "datos/backups/b.xlsx", "datos/ultimo_reset.json"):
        _escribir(app / rel, f"dato {rel}")
    cfg = Config(base=base, planilla=pl, excel_motor="openpyxl", modo_real=True)
    return app, cfg


def _foto(raiz: Path, sin=()) -> dict:
    return {f.relative_to(raiz).as_posix(): sha256(f) for f in raiz.rglob("*")
            if f.is_file() and f.relative_to(raiz).as_posix() not in sin}


def _real(tmp_path, monkeypatch):
    raiz = tmp_path / OD
    app, cfg = _arbol(tmp_path, monkeypatch, raiz)
    docs = tmp_path / "Documents"
    docs.mkdir()
    return raiz, app, cfg, docs


REGISTROS = ("datos/borrado.json", "datos/borrado_prueba.json")


def _senuelos_intactos(raiz, app, antes_od, antes_app):
    for s in SENUELOS_OD:
        assert sha256(raiz / s) == antes_od[s], s
    for s in SENUELOS_APP:
        assert sha256(app / s) == antes_app[s], s


def test_borrar_y_deshacer_deja_todo_igual(tmp_path, monkeypatch):
    raiz, app, cfg, docs = _real(tmp_path, monkeypatch)
    antes_od, antes_app = _foto(raiz), _foto(app)

    inv = borrado.inventario(cfg, "real")
    assert inv["pdf"] == 3 and inv["carpetas_proveedores"] == 2 and inv["en_entrada"] == 3
    assert inv["filas"]["Vehicular"] == 1 and inv["filas"]["Hogar"] == 2 and inv["proveedores"] == 1
    assert {p.name for p in inv["quedan"]} == {"mio.docx", "Leeme.txt"}

    reg = borrado.borrar(cfg, "real", carpeta_zip=docs, log=lambda s: None)
    assert reg["estado"] == "hecho"
    _senuelos_intactos(raiz, app, antes_od, antes_app)
    # gone from their places, kept in _Borrado
    b_od = Path(reg["carpetas"][0])
    assert b_od.parent == raiz and b_od.name.startswith("_Borrado ")
    assert not (raiz / "PROVEEDORES RUA.xlsx").exists() and not (raiz / "Planillas exportadas" / "Cierre octubre.xlsx").exists()
    assert not (raiz / "DEUDA A PAGAR ARG - respaldo 2026-09-01.xlsx").exists()
    assert (b_od / "Facturas pendientes" / "Vehicular" / "Remolques Inventados" / "Factura 0001-00000001.pdf").is_file()
    assert not (app / "datos" / "lotes" / "L1").exists() and not (app / "datos" / "tandas.json").exists()
    # Setup again: empty structure, csv with the header only
    csv = raiz / "Facturas pendientes" / "_Script" / "proveedores.csv"
    assert csv.read_text(encoding="utf-8-sig", newline="") == preparar.CSV_ENCABEZADO
    assert (raiz / "Facturas pendientes" / "_Entrada" / "Hogar").is_dir()
    assert not list((raiz / "Facturas pendientes" / "Vehicular").iterdir())
    # workbook emptied in place, column I and layout kept; its copy in _Borrado
    wb = load_workbook(raiz / "DEUDA A PAGAR ARG.xlsx")
    assert wb["Vehicular"]["A12"].value is None and wb["Hogar"]["A13"].value is None
    assert wb["Proveedores"]["B5"].value is None
    assert wb["Vehicular"]["I12"].value == '=IF(G12="","",G12+H12)' and wb["Vehicular"]["A11"].value == "Proveedor"
    assert sha256(b_od / "DEUDA A PAGAR ARG.xlsx") == antes_od["DEUDA A PAGAR ARG.xlsx"]
    # verified zip outside OneDrive and the app
    z = Path(reg["zip"])
    assert z.parent == docs and borrado.verificar_zip(z) == len(inv["archivos"]) + 2   # + workbook + config.toml

    reg = borrado.deshacer(cfg, "real", log=lambda s: None)
    assert reg["estado"] == "deshecho" and reg["avisos"] == []
    assert _foto(raiz) == antes_od
    assert _foto(app, sin=REGISTROS) == antes_app
    assert not any(p.name.startswith("_Borrado") for p in list(raiz.iterdir()) + list(app.iterdir()))


def test_vaciar_solo_manda_las_carpetas_borrado(tmp_path, monkeypatch):
    raiz, app, cfg, docs = _real(tmp_path, monkeypatch)
    antes_od, antes_app = _foto(raiz), _foto(app)
    reg = borrado.borrar(cfg, "real", carpeta_zip=docs, log=lambda s: None)
    papelera = []
    monkeypatch.setattr(borrado, "_a_la_papelera", lambda p: (papelera.append(p), os.rename(p, tmp_path / f"papelera{len(papelera)}")))
    with pytest.raises(borrado.BorradoError):
        borrado.vaciar(cfg, "real", "vaciar", log=lambda s: None)
    assert papelera == []
    reg = borrado.vaciar(cfg, "real", "VACIAR", log=lambda s: None)
    assert reg["estado"] == "vaciado" and [str(p) for p in papelera] == reg["carpetas"]
    _senuelos_intactos(raiz, app, antes_od, antes_app)
    with pytest.raises(borrado.BorradoError):
        borrado.deshacer(cfg, "real", log=lambda s: None)


def test_fallo_a_mitad_devuelve_todo(tmp_path, monkeypatch):
    raiz, app, cfg, docs = _real(tmp_path, monkeypatch)
    antes_od, antes_app = _foto(raiz), _foto(app)
    original, n = borrado._mover, []

    def mover(src, dst):
        n.append(src)
        if len(n) == 4:
            raise OSError("archivo tomado por OneDrive")
        original(src, dst)
    monkeypatch.setattr(borrado, "_mover", mover)
    with pytest.raises(borrado.BorradoError, match="devolví todo"):
        borrado.borrar(cfg, "real", carpeta_zip=docs, log=lambda s: None)
    assert _foto(raiz) == antes_od and _foto(app, sin=REGISTROS) == antes_app
    assert borrado.estado("real")["estado"] == "revertido"
    assert not any(p.name.startswith("_Borrado") for p in list(raiz.iterdir()) + list(app.iterdir()))


def test_fallo_en_la_planilla_devuelve_todo(tmp_path, monkeypatch):
    raiz, app, cfg, docs = _real(tmp_path, monkeypatch)
    antes_od, antes_app = _foto(raiz), _foto(app)

    def falla(libro, prov, log):
        libro.hoja("Vehicular").poner(12, "A", "a medias")
        libro.guardar()
        raise RuntimeError("Excel se colgó")
    monkeypatch.setattr(preparar, "vaciar_libro", falla)
    with pytest.raises(borrado.BorradoError):
        borrado.borrar(cfg, "real", carpeta_zip=docs, log=lambda s: None)
    assert _foto(raiz) == antes_od and _foto(app, sin=REGISTROS) == antes_app


def test_deshacer_no_pisa_lo_nuevo(tmp_path, monkeypatch):
    raiz, app, cfg, docs = _real(tmp_path, monkeypatch)
    reg = borrado.borrar(cfg, "real", carpeta_zip=docs, log=lambda s: None)
    nueva = _escribir(raiz / "Facturas pendientes" / "Vehicular" / "Remolques Inventados" / "Factura 0001-00000001.pdf",
                      "cargada después del borrado")
    csv = raiz / "Facturas pendientes" / "_Script" / "proveedores.csv"
    reg = borrado.deshacer(cfg, "real", log=lambda s: None)
    assert nueva.read_text(encoding="utf-8") == "cargada después del borrado"
    assert reg["estado"] == "deshecho a medias" and any("no lo piso" in a for a in reg["avisos"])
    vieja = Path(reg["carpetas"][0]) / "Facturas pendientes" / "Vehicular" / "Remolques Inventados" / "Factura 0001-00000001.pdf"
    assert vieja.read_text(encoding="utf-8").startswith("contenido de")
    assert (nueva.parent / "Factura 0001-00000002.pdf").is_file()               # the rest came back
    assert csv.read_bytes() == CSV.encode("utf-8")                            # the real csv, not the empty one


def test_planilla_cambiada_despues_no_se_pisa(tmp_path, monkeypatch):
    raiz, app, cfg, docs = _real(tmp_path, monkeypatch)
    borrado.borrar(cfg, "real", carpeta_zip=docs, log=lambda s: None)
    pl = raiz / "DEUDA A PAGAR ARG.xlsx"
    wb = load_workbook(pl)
    wb["Hogar"]["A12"] = "cargada después"
    wb.save(pl)
    reg = borrado.deshacer(cfg, "real", log=lambda s: None)
    assert load_workbook(pl)["Hogar"]["A12"].value == "cargada después"
    assert any("planilla cambió" in a for a in reg["avisos"]) and Path(reg["planilla"]["copia"]).is_file()


def test_practicar_solo_toca_la_carpeta_de_prueba(tmp_path, monkeypatch):
    raiz_real = tmp_path / OD
    app, cfg_real = _arbol(tmp_path, monkeypatch, raiz_real)
    prueba = app / "prueba" / OD
    _, cfg = _arbol(tmp_path / "x", monkeypatch, prueba)
    monkeypatch.setenv("APPFACTURAS_HOME", str(app))
    cfg.modo_real = False
    docs = tmp_path / "Documents"
    docs.mkdir()
    antes_real, antes_app = _foto(raiz_real), _foto(app)
    reg = borrado.borrar(cfg, "prueba", carpeta_zip=docs, log=lambda s: None)
    assert reg["carpetas"] == [str(prueba / Path(reg["carpetas"][0]).name)]
    assert _foto(raiz_real) == antes_real
    despues = _foto(app, sin=REGISTROS)
    assert {k: v for k, v in despues.items() if not k.startswith("prueba/")} == \
        {k: v for k, v in antes_app.items() if not k.startswith("prueba/")}           # datos\ untouched
    assert not (prueba / "PROVEEDORES RUA.xlsx").exists()
    borrado.deshacer(cfg, "prueba", log=lambda s: None)
    assert _foto(app, sin=REGISTROS) == antes_app


def test_frena_antes_de_mover_nada(tmp_path, monkeypatch):
    raiz, app, cfg, docs = _real(tmp_path, monkeypatch)
    antes_od, antes_app = _foto(raiz), _foto(app)
    sin_cambios = lambda: _foto(raiz) == antes_od and _foto(app, sin=REGISTROS) == antes_app and not list(docs.iterdir())

    with pytest.raises(borrado.BorradoError, match="Practicar"):         # practice mode on the real OneDrive
        borrado.borrar(cfg, "prueba", carpeta_zip=docs, log=lambda s: None)
    cfg.modo_real = False
    with pytest.raises(borrado.BorradoError, match="Solo simular"):
        borrado.borrar(cfg, "real", carpeta_zip=docs, log=lambda s: None)
    cfg.modo_real = True
    monkeypatch.setattr(borrado, "_tomado", lambda p: "está abierto en otro programa" if p.suffix == ".pdf" else "")
    with pytest.raises(borrado.BorradoError, match="abierto en otro programa"):
        borrado.borrar(cfg, "real", carpeta_zip=docs, log=lambda s: None)
    monkeypatch.setattr(borrado, "_tomado", lambda p: "")
    monkeypatch.setattr(borrado, "_planilla_sin_guardar", lambda c, p: True)
    with pytest.raises(borrado.BorradoError, match="sin guardar"):
        borrado.borrar(cfg, "real", carpeta_zip=docs, log=lambda s: None)
    monkeypatch.setattr(borrado, "_planilla_sin_guardar", lambda c, p: False)
    with pytest.raises(borrado.BorradoError, match="ni en la app"):      # backup inside the OneDrive
        borrado.borrar(cfg, "real", carpeta_zip=raiz / "Escritorio", log=lambda s: None)
    otra = Config(base=tmp_path / "Facturas pendientes", planilla=raiz / "DEUDA A PAGAR ARG.xlsx",
                  excel_motor="openpyxl", modo_real=True)
    with pytest.raises(borrado.BorradoError, match="juntas"):
        borrado.borrar(otra, "real", carpeta_zip=docs, log=lambda s: None)
    assert sin_cambios()

    borrado.borrar(cfg, "real", carpeta_zip=docs, log=lambda s: None)
    with pytest.raises(borrado.BorradoError, match="sin deshacer"):      # a second one before undoing the first
        borrado.borrar(cfg, "real", carpeta_zip=docs, log=lambda s: None)


def test_link_adentro_frena(tmp_path, monkeypatch):
    raiz, app, cfg, docs = _real(tmp_path, monkeypatch)
    afuera = tmp_path / "afuera"
    _escribir(afuera / "importante.txt", "no es de la app")
    link = raiz / "Facturas pendientes" / "Hogar" / "atajo"
    try:
        import _winapi
        _winapi.CreateJunction(str(afuera), str(link))
    except (ImportError, OSError):
        pytest.skip("no se puede crear una junction acá")
    with pytest.raises(borrado.BorradoError, match="link"):
        borrado.inventario(cfg, "real")
    assert (afuera / "importante.txt").is_file()
