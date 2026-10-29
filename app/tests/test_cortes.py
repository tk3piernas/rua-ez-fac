"""What a cut (app closed, PC off) or a shortcut around the app could break (made-up data only)."""
import json
from pathlib import Path

import pytest

from appfacturas import borrado, claude_cowork, cli, ejecutor, util
from appfacturas.util import sha256


def test_la_cli_no_mueve_ni_escribe(tmp_path, monkeypatch):
    """procesar / deshacer / excel-pendiente skipped the work mode (they could touch the real OneDrive)."""
    monkeypatch.setenv("APPFACTURAS_HOME", str(tmp_path / "app"))
    for cmd in (["procesar", "--real", "--si"], ["deshacer", "--si"], ["excel-pendiente"]):
        with pytest.raises(SystemExit):
            cli.main(cmd)
    base = tmp_path / "OneDrive - Rua Asistencia" / "Facturas pendientes"
    pdf = base / "_Entrada" / "Hogar" / "suelta.pdf"
    pdf.parent.mkdir(parents=True)
    pdf.write_bytes(b"%PDF-1.4 inventado")
    (base / "_Script").mkdir()
    (base / "_Script" / "proveedores.csv").write_text("CUIT;Area;Proveedor;Carpeta\r\n", encoding="utf-8")
    assert cli.main(["simular", "--entrada", "--base", str(base), "--sin-planilla"]) == 0
    assert pdf.read_bytes() == b"%PDF-1.4 inventado"                       # simular never moves


def test_guardar_json_sobrevive_un_corte(tmp_path, monkeypatch):
    r = tmp_path / "diario.json"
    util.guardar_json(r, [{"op": "mover"}])

    def corta(*a, **kw):
        a[1].write('[{"op": "mov')                                         # half written, then the cut
        raise KeyboardInterrupt
    monkeypatch.setattr(json, "dump", corta)
    with pytest.raises(KeyboardInterrupt):
        util.guardar_json(r, [{"op": "mover"}, {"op": "copiar"}])
    assert json.loads(r.read_text(encoding="utf-8")) == [{"op": "mover"}]   # the old one is intact


def test_un_pdf_movido_justo_antes_de_cortarse_vuelve(tmp_path):
    """The file moved but the app closed before journaling it: undo still brings it back."""
    entrada, destino = tmp_path / "_Entrada" / "Hogar" / "f.pdf", tmp_path / "Hogar" / "Prov" / "Factura 1.pdf"
    destino.parent.mkdir(parents=True)
    destino.write_bytes(b"factura inventada")
    entrada.parent.mkdir(parents=True)
    lote = tmp_path / "lote"
    d = ejecutor.Diario(lote)
    d.anotar({"op": "moviendo", "como": "mover", "de": str(entrada), "a": str(destino), "sha256": sha256(destino)})
    problemas = ejecutor.deshacer(lote, log=lambda s: None)
    assert problemas == [] and entrada.read_bytes() == b"factura inventada" and not destino.exists()


def test_un_movimiento_que_no_llego_a_pasar_no_se_inventa(tmp_path):
    entrada, destino = tmp_path / "_Entrada" / "f.pdf", tmp_path / "Prov" / "Factura 1.pdf"
    entrada.parent.mkdir(parents=True)
    entrada.write_bytes(b"x")
    d = ejecutor.Diario(tmp_path / "lote")
    d.anotar({"op": "moviendo", "como": "mover", "de": str(entrada), "a": str(destino), "sha256": sha256(entrada)})
    assert d.completar_cortados() == 0 and not any(p["op"] == "mover" for p in d.pasos)


def test_borrar_un_pedido_a_claude_va_a_la_papelera(tmp_path, monkeypatch):
    papelera = []
    monkeypatch.setattr(borrado, "_a_la_papelera", papelera.append)
    pedido = tmp_path / "Facturas pendientes" / claude_cowork.CARPETA / "2026-10-28_120000"
    pedido.mkdir(parents=True)
    (pedido / "pedido.json").write_text("{}", encoding="utf-8")
    claude_cowork.borrar(pedido)
    assert papelera == [pedido] and pedido.exists()                         # moved by the recycle bin, not rmtree
    otra = tmp_path / "Facturas pendientes" / "Vehicular"
    otra.mkdir()
    with pytest.raises(ValueError):
        claude_cowork.borrar(otra)


def _planilla_cortada(tmp_path, guardada: bool):
    """A batch whose Excel run was cut: the journal has the row but neither 'excel_guardado' nor 'excel_revertido'."""
    from openpyxl import load_workbook
    from appfacturas import preparar
    from appfacturas.config import Config
    pl = tmp_path / "OneDrive - Rua Asistencia" / "DEUDA A PAGAR ARG.xlsx"
    preparar.crear_planilla_vacia(pl)
    link = '=HYPERLINK("https://ejemplo.invalid/f.pdf","Factura 0001-00000001")'
    if guardada:
        wb = load_workbook(pl)
        wb["Vehicular"]["A12"], wb["Vehicular"]["J12"] = "Remolques Inventados", link
        wb.save(pl)
    lote = tmp_path / "lote"
    d = ejecutor.Diario(lote)
    d.anotar({"op": "respaldo_excel", "ruta": "x"})
    d.anotar({"op": "fila_excel", "libro": str(pl), "hoja": "Vehicular", "fila": 12,
              "antes": {c: "" for c in "ABCDEFGHIJ"}, "escrito": {"A": "Remolques Inventados", "J": link},
              "colores": {}, "comentario_antes": None, "comentado": False, "factura": "f1"})
    return Config(base=pl.parent / "Facturas pendientes", planilla=pl, excel_motor="openpyxl"), d, lote, pl


@pytest.mark.parametrize("guardada", [True, False])
def test_escritura_cortada_se_decide_mirando_la_planilla(tmp_path, guardada):
    from appfacturas import excel
    cfg, d, _, _ = _planilla_cortada(tmp_path, guardada)
    assert excel.reconocer_cortada(cfg, d, log=lambda s: None) is guardada
    assert bool(excel.pasos_guardados(d.pasos)) is guardada                 # saved rows are never written twice
    assert excel.tanda_cortada(d.pasos) == [] and excel.reconocer_cortada(cfg, d) is None


def test_deshacer_saca_las_filas_de_una_escritura_cortada(tmp_path):
    from openpyxl import load_workbook
    cfg, _, lote, pl = _planilla_cortada(tmp_path, guardada=True)
    assert ejecutor.deshacer(lote, log=lambda s: None, cfg=cfg) == []
    ws = load_workbook(pl)["Vehicular"]
    assert ws["A12"].value is None and ws["J12"].value is None


def test_terminar_excel_no_toca_lo_que_no_es_excel():
    import os
    from appfacturas import excel
    assert excel._terminar_excel(os.getpid()) is False                     # this python is still alive
