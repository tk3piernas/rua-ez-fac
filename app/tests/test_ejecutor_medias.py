"""Batches cut halfway (A3), files that did not move (A7), where a batch was loaded (A11). Synthetic invoices."""
import pytest
from openpyxl import load_workbook

from appfacturas import ejecutor, excel
from appfacturas.lote import Armador, entradas_de_carpeta
from test_lote_sintetico import HOY, factura, sandbox  # noqa: F401  (fixture)


def _armar(cfg, tmp_path):
    trabajo = tmp_path / "datos" / "lotes" / "2026-10-27_100000"
    ents, txts = entradas_de_carpeta(cfg)
    return Armador(cfg, simular=False, hoy=HOY, trabajo=trabajo).procesar(ents, txts), trabajo


def _filas_escritas(cfg, hoja="Vehicular"):
    ws = load_workbook(cfg.planilla)[hoja]
    return [ws[f"A{r}"].value for r in range(excel.PRIMERA, excel.ULTIMA + 1) if ws[f"A{r}"].value]


def _dos_facturas(cfg):
    factura(cfg.entrada / "Vehicular" / "uno.pdf", 31)
    factura(cfg.entrada / "Vehicular" / "dos.pdf", 32)


def test_lote_cortado_a_medias_se_termina_y_se_deshace(sandbox, tmp_path, monkeypatch):
    _dos_facturas(sandbox)
    lote, trabajo = _armar(sandbox, tmp_path)

    def corte(*a, **kw):
        raise KeyboardInterrupt          # the app closed while Excel was writing
    monkeypatch.setattr(ejecutor, "escribir_excel", corte)
    with pytest.raises(KeyboardInterrupt):
        ejecutor.Ejecutor(sandbox, lote, trabajo, log=lambda s: None, modo="prueba").ejecutar()
    monkeypatch.undo()
    assert (trabajo / "lote_confirmado.json").is_file()
    assert ejecutor.etiqueta(trabajo) == "quedó a medias"
    assert ejecutor.ultimo_confirmado(trabajo.parent) == trabajo
    r = ejecutor.continuar(sandbox, trabajo, log=lambda s: None)
    assert r["excel"] == "escrito" and len(_filas_escritas(sandbox)) == 2
    assert ejecutor.estado(trabajo)["estado"] == "confirmado"
    # continuing again does not duplicate anything
    ejecutor.continuar(sandbox, trabajo, log=lambda s: None)
    assert len(_filas_escritas(sandbox)) == 2
    assert ejecutor.deshacer(trabajo, log=lambda s: None, cfg=sandbox) == []
    assert _filas_escritas(sandbox) == []
    assert sorted(p.name for p in (sandbox.entrada / "Vehicular").iterdir()) == ["dos.pdf", "uno.pdf"]


def test_archivo_que_no_se_movio_no_escribe_su_fila(sandbox, tmp_path, monkeypatch):
    _dos_facturas(sandbox)
    lote, trabajo = _armar(sandbox, tmp_path)
    mover = ejecutor._mover

    def bloqueado(de, a, intentos=4):
        if de.name == "dos.pdf":
            raise PermissionError("OneDrive lo tiene tomado")
        mover(de, a, intentos=1)
    monkeypatch.setattr(ejecutor, "_mover", bloqueado)
    res = ejecutor.Ejecutor(sandbox, lote, trabajo, log=lambda s: None).ejecutar()
    est = ejecutor.estado(trabajo)
    assert res["sin_archivo"] == 1 and len(est["archivos_pendientes"]) == 1
    assert est["estado"] == "confirmado con errores"
    assert len(_filas_escritas(sandbox)) == 1                         # only the one whose PDF moved
    assert (sandbox.entrada / "Vehicular" / "dos.pdf").exists()
    monkeypatch.undo()
    r = ejecutor.continuar(sandbox, trabajo, log=lambda s: None)
    assert r["errores"] == [] and r["sin_archivo"] == 0
    assert len(_filas_escritas(sandbox)) == 2
    assert ejecutor.estado(trabajo)["archivos_pendientes"] == []
    assert ejecutor.deshacer(trabajo, log=lambda s: None, cfg=sandbox) == []
    assert _filas_escritas(sandbox) == []


def test_estado_dice_donde_se_cargo(sandbox, tmp_path):
    factura(sandbox.entrada / "Vehicular" / "uno.pdf", 33)
    lote, trabajo = _armar(sandbox, tmp_path)
    ejecutor.Ejecutor(sandbox, lote, trabajo, log=lambda s: None, modo="prueba").ejecutar()
    e = ejecutor.estado(trabajo)
    assert (e["modo"], e["planilla"], e["base"]) == ("prueba", str(sandbox.planilla), str(sandbox.base))
    assert ejecutor.donde(trabajo) == ("prueba", str(sandbox.planilla))


def test_memoria_que_no_se_pudo_deshacer_queda_pendiente(sandbox, tmp_path, monkeypatch):
    from appfacturas import maestro
    _dos_facturas(sandbox)
    lote, trabajo = _armar(sandbox, tmp_path)
    ejecutor.Ejecutor(sandbox, lote, trabajo, log=lambda s: None).ejecutar()
    (trabajo / ejecutor.HIST_MEMORIA).write_text("[]", encoding="utf-8")
    llamadas = []

    def falla(*a, **kw):
        llamadas.append(1)
        if len(llamadas) == 1:
            raise excel.LibroConCambios("PROVEEDORES RUA.xlsx está abierta con cambios sin guardar")
        return []
    monkeypatch.setattr(maestro, "deshacer", falla)
    problemas = ejecutor.deshacer(trabajo, log=lambda s: None, cfg=sandbox)
    assert any("quedó pendiente" in p for p in problemas)
    assert ejecutor.estado(trabajo)["memoria_pendiente"] is True
    assert ejecutor.deshacer(trabajo, log=lambda s: None, cfg=sandbox) == []     # retries only the memory
    assert ejecutor.estado(trabajo)["memoria_pendiente"] is False
    with pytest.raises(ejecutor.YaDeshecho):
        ejecutor.deshacer(trabajo, log=lambda s: None, cfg=sandbox)
