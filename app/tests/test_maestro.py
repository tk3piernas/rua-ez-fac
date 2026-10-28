"""Provider memory (PROVEEDORES RUA.xlsx): read pasted text, save, update and undo. Invented numbers only."""
from openpyxl import load_workbook

from appfacturas import maestro
from appfacturas.config import Config
from appfacturas.util import cbu_valido, cuit_valido


def _cuit(base10: str) -> str:
    return next(base10 + str(d) for d in range(10) if cuit_valido(base10 + str(d)))


def _cbu(banco7: str, cuenta13: str) -> str:
    b = next(banco7 + str(d) for d in range(10) if cbu_valido(banco7 + str(d) + "0" * 14))
    return next(b + cuenta13 + str(d) for d in range(10) if cbu_valido(b + cuenta13 + str(d)))


CUIT = _cuit("3071234567")
CBU = _cbu("0170099", "2200001234567")


def _cfg(tmp_path):
    raiz = tmp_path / "OneDrive - Rua Asistencia"
    return Config(base=raiz / "Facturas pendientes", planilla=raiz / "DEUDA A PAGAR ARG.xlsx", excel_motor="openpyxl")


TEXTO = f"""Hola! Te paso los datos
Razón social: Grúas del Sur SRL
CUIT: {CUIT[:2]}-{CUIT[2:10]}-{CUIT[10]}
Teléfono: 11 5555-0000
mail: pagos@gruasdelsur.example

Banco: BBVA
CBU: {CBU}
Alias: GRUAS.SUR.PAGOS
"""


def test_lee_texto_pegado():
    vals, avisos = maestro.desde_texto(TEXTO, "Vehicular")
    assert vals[0] == "Vehicular" and vals[1] == "Grúas del Sur SRL" and vals[2] == CUIT
    assert vals[7] == "pagos@gruasdelsur.example" and vals[11] == CBU and vals[12] == "GRUAS.SUR.PAGOS"
    assert all(ok for ok, _ in maestro.chequeos(vals))


def test_guarda_actualiza_y_deshace(tmp_path):
    cfg = _cfg(tmp_path)
    hist = tmp_path / "hist.json"
    vals, _ = maestro.desde_texto(TEXTO, "Vehicular")
    assert len(maestro.guardar(cfg, [(vals, False)], hist, "test", log=lambda s: None)) == 1
    r = maestro.ruta(cfg)
    assert r.is_file() and r.parent == cfg.planilla.parent
    mem = maestro.leer(cfg)
    assert mem[CUIT]["valores"][11] == CBU and mem[CUIT]["fila"] == 2
    assert load_workbook(r)["Proveedores"]["L2"].number_format == "@"           # CBU as text

    # same provider again: same row; blanks never erase; solo_vacios only fills blanks
    otro = list(vals)
    otro[5], otro[11], otro[7] = "Ana", "", "otro@x.example"
    maestro.guardar(cfg, [(otro, True)], hist, "test", log=lambda s: None)
    m = maestro.leer(cfg)[CUIT]
    assert m["fila"] == 2 and m["valores"][5] == "Ana" and m["valores"][11] == CBU
    assert m["valores"][7] == "pagos@gruasdelsur.example"
    assert maestro.guardar(cfg, [(otro, True)], hist, "test", log=lambda s: None) == []   # nothing new

    assert maestro.deshacer(cfg, hist, log=lambda s: None) == []
    assert maestro.leer(cfg)[CUIT]["valores"][5] == ""
    # the first save created the file: undoing it removes the file
    assert maestro.deshacer(cfg, hist, log=lambda s: None) == [] and not r.exists()


def test_deshacer_respeta_cambios_posteriores(tmp_path):
    cfg = _cfg(tmp_path)
    maestro.crear(maestro.ruta(cfg))
    hist = tmp_path / "hist.json"
    vals, _ = maestro.desde_texto(TEXTO)
    maestro.guardar(cfg, [(vals, False)], hist, "test", log=lambda s: None)
    wb = load_workbook(maestro.ruta(cfg))
    wb["Proveedores"]["F2"] = "lo cambió una persona"
    wb.save(maestro.ruta(cfg))
    problemas = maestro.deshacer(cfg, hist, log=lambda s: None)
    assert problemas and "cambió después" in problemas[0]
    assert maestro.leer(cfg)[CUIT]["valores"][5] == "lo cambió una persona"
