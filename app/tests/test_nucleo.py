"""F5: core cases without a test until now. Synthetic invoices, invented provider."""
import sys
import types
import zipfile

from openpyxl import load_workbook

from appfacturas import config, ejecutor, excel
from appfacturas.lote import aplicar_exclusiones
from appfacturas.model import Estado
from test_ejecutor_medias import _armar
from test_lote_sintetico import factura, sandbox  # noqa: F401  (fixture)


def _confirmar_uno(cfg, tmp_path, con_excel=True):
    factura(cfg.entrada / "Vehicular" / "uno.pdf", 101)
    lote, trabajo = _armar(cfg, tmp_path)
    ejecutor.Ejecutor(cfg, lote, trabajo, log=lambda s: None, con_excel=con_excel).ejecutar()
    return trabajo


def test_deshacer_no_toca_filas_cambiadas_a_mano(sandbox, tmp_path):
    for col, valor in (("A", "Otro proveedor"), ("J", '=HYPERLINK("https://x/otro.pdf","Otra")')):
        trabajo = _confirmar_uno(sandbox, tmp_path / col)
        wb = load_workbook(sandbox.planilla)
        wb["Vehicular"][f"{col}{excel.PRIMERA}"] = valor
        wb.save(sandbox.planilla)
        problemas = ejecutor.deshacer(trabajo, log=lambda s: None, cfg=sandbox)
        assert any("cambió después" in p for p in problemas), col
        assert load_workbook(sandbox.planilla)["Vehicular"][f"{col}{excel.PRIMERA}"].value == valor
        wb = load_workbook(sandbox.planilla)                     # clean for the next round
        for c in "ABCDEFGJ":
            wb["Vehicular"][f"{c}{excel.PRIMERA}"] = None
        wb.save(sandbox.planilla)


def test_reintentar_con_la_hoja_faltante(sandbox, tmp_path):
    trabajo = _confirmar_uno(sandbox, tmp_path, con_excel=False)
    assert ejecutor.estado(trabajo)["excel"] == "pendiente"
    wb = load_workbook(sandbox.planilla)
    del wb["Vehicular"]
    wb.save(sandbox.planilla)
    estado, msg = ejecutor.reintentar_excel(sandbox, trabajo, log=lambda s: None)
    assert estado == "pendiente" and "Vehicular" in msg
    assert ejecutor.estado(trabajo)["excel_pendiente"] == 2            # the row and the new provider


def test_exclusion_parcial_de_un_zip(sandbox, tmp_path):
    a, b = tmp_path / "a.pdf", tmp_path / "b.pdf"
    factura(a, 102)
    factura(b, 103)
    with zipfile.ZipFile(sandbox.entrada / "Vehicular" / "dos.zip", "w") as z:
        z.write(a, "a.pdf")
        z.write(b, "b.pdf")
    lote, _ = _armar(sandbox, tmp_path)
    fuera = next(f for f in lote.facturas if f.ext.nro == 103)
    for r in lote.filas:
        if r.factura_id == fuera.id:
            r.excluida = True
    assert aplicar_exclusiones(lote) == {fuera.id}
    assert fuera.estado == Estado.EXCLUIDA
    copias = [p for p in lote.plan if p["op"] == "copiar"]
    assert copias and all(p["factura"] != fuera.id for p in copias)
    assert not [p for p in lote.plan if p["op"] == "mover_zip"]          # the zip stays in _Entrada


def test_url_desde_onedrive_con_registro_simulado(tmp_path, monkeypatch):
    base = tmp_path / "OneDrive - Empresa" / "Facturas pendientes"
    clave = r"Software\Microsoft\OneDrive\Accounts\Business1"
    claves = {clave: {"UserFolder": str(base.parent), "UserEmail": "ana.perez@empresa.com",
                      "ServiceEndpointUri": "https://empresa-my.sharepoint.com/personal/ana_perez_empresa_com/_api"}}

    class Clave:
        def __init__(self, ruta):
            if ruta not in claves:
                raise OSError(ruta)
            self.v = claves[ruta]

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    winreg = types.SimpleNamespace(HKEY_CURRENT_USER=0, OpenKey=lambda h, ruta: Clave(ruta),
                                   QueryValueEx=lambda k, n: (k.v[n], 1))
    monkeypatch.setitem(sys.modules, "winreg", winreg)
    monkeypatch.setattr(config.os, "name", "nt")
    esperada = "https://empresa-my.sharepoint.com/personal/ana_perez_empresa_com/Documents/Facturas pendientes"
    assert config.url_desde_onedrive(base) == esperada
    # without the endpoint it needs the tenant from config.toml
    claves[clave]["ServiceEndpointUri"] = ""
    assert config.url_desde_onedrive(base) == ""
    assert config.url_desde_onedrive(base, "https://empresa-my.sharepoint.com/") == esperada
