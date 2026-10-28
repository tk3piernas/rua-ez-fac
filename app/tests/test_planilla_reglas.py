"""B2: only H and J can be formulas. B4: CUIT typed by hand as a number in the Proveedores sheet."""
from openpyxl import load_workbook

from appfacturas import excel, planilla, preparar
from appfacturas.util import cuit_valido

CUIT = next("3070000222" + str(d) for d in range(10) if cuit_valido("3070000222" + str(d)))


def test_texto_con_igual_no_es_formula(tmp_path):
    ruta = tmp_path / "DEUDA A PAGAR ARG.xlsx"
    preparar.crear_planilla_vacia(ruta)
    libro = excel.LibroOpenpyxl(ruta)
    h = libro.hoja("Proveedores")
    h.poner(5, "O", "=ver mail")
    h.poner(5, "J", '=HYPERLINK("https://x/a.pdf","Factura")', formula=True)
    libro.guardar()
    ws = load_workbook(ruta)["Proveedores"]
    assert ws["O5"].data_type == "s" and ws["O5"].value == "=ver mail"
    assert ws["J5"].data_type == "f"


def test_cuit_cargado_como_numero(tmp_path):
    ruta = tmp_path / "DEUDA A PAGAR ARG.xlsx"
    preparar.crear_planilla_vacia(ruta)
    wb = load_workbook(ruta)
    ws = wb["Proveedores"]
    ws["B5"], ws["C5"], ws["L5"] = "Prov Inventado", float(CUIT), 1.4e21
    wb.save(ruta)
    (p,) = planilla.leer(ruta).proveedores
    assert p["cuit"] == CUIT and p["cbu"] == ""
