"""Setup and workbook reset (no real invoices needed: this file also goes to the repo)."""
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

from appfacturas import planilla, preparar
from appfacturas.config import Config


def _cfg(tmp_path):
    base = tmp_path / "OneDrive - Rua Asistencia" / "Facturas pendientes"
    return Config(base=base, planilla=base.parent / "DEUDA A PAGAR ARG.xlsx", excel_motor="openpyxl",
                  url_base="https://rua.example/Documents/Facturas pendientes")


def test_setup_crea_todo_y_no_pisa(tmp_path, monkeypatch):
    monkeypatch.setenv("APPFACTURAS_HOME", str(tmp_path / "app"))
    cfg = _cfg(tmp_path)
    cfg.base.parent.mkdir(parents=True)
    plan = preparar.plan_setup(cfg)
    assert cfg.planilla in plan["archivos"] and cfg.base in plan["dirs"]
    hecho = preparar.ejecutar_setup(cfg, plan)
    assert hecho
    for s in preparar.SUBCARPETAS:
        assert (cfg.base / s).is_dir()
    csv = cfg.base / "_Script" / "proveedores.csv"
    assert csv.read_text(encoding="utf-8-sig").startswith("CUIT;Area;Proveedor;Carpeta")
    assert (tmp_path / "app" / "datos" / "lotes").is_dir()
    pl = planilla.leer(cfg.planilla)
    assert pl.leida and pl.filas == [] and float(pl.alicuota["Vehicular"]) == 0.21
    # second run: nothing to do, nothing overwritten
    csv.write_text("CUIT;Area;Proveedor;Carpeta\r\n1;Hogar;X;X\r\n", encoding="utf-8-sig")
    plan2 = preparar.plan_setup(cfg)
    assert plan2["dirs"] == [] and plan2["archivos"] == []
    assert "1;Hogar;X" in csv.read_text(encoding="utf-8-sig")


def test_reset_vacia_en_el_lugar_y_respalda(tmp_path):
    cfg = _cfg(tmp_path)
    preparar.crear_planilla_vacia(cfg.planilla)
    wb = load_workbook(cfg.planilla)
    ws = wb["Vehicular"]
    ws["A12"], ws["B12"], ws["G12"], ws["H12"] = "Remolques X", "412345", 1000, 0        # H overwritten (Factura C)
    ws["J12"] = '=HYPERLINK("https://x/a.pdf","Factura 0001-00000001")'
    ws["A13"], ws["G13"] = "Otro", 500
    ws["J13"].hyperlink = "https://x/b.pdf"
    ws["A13"].comment = Comment("posible duplicado", "x")
    for c in "ABCDEFGHIJ":
        ws[f"{c}13"].fill = PatternFill("solid", fgColor="FFFFF2CC")
    wb["Proveedores"]["A5"], wb["Proveedores"]["B5"] = "Vehicular", "Remolques X"
    wb.save(cfg.planilla)

    res = preparar.reset_planilla(cfg, tmp_path / "backups", vaciar_proveedores=False, log=lambda s: None)
    assert res["filas"] == {"Vehicular": 2, "Hogar": 0, "Persona": 0, "ServiciosYa": 0}
    assert len(res["respaldos"]) == 2 and all(__import__("pathlib").Path(p).is_file() for p in res["respaldos"])

    wb = load_workbook(cfg.planilla)
    ws = wb["Vehicular"]
    for c in "ABCDEFGJ":
        assert ws[f"{c}12"].value is None and ws[f"{c}13"].value is None
    assert ws["H12"].value == '=IF(G12="","",ROUND(G12*$B$4,2))'      # formula back
    assert ws["I12"].value == '=IF(G12="","",G12+H12)'                 # I untouched
    assert ws["A13"].comment is None and ws["J13"].hyperlink is None
    assert ws["A13"].fill.fill_type is None
    assert ws["B4"].value == 0.21 and ws["A11"].value == "Proveedor"   # layout kept
    assert wb["Proveedores"]["B5"].value == "Remolques X"              # kept unless asked

    preparar.reset_planilla(cfg, tmp_path / "backups", vaciar_proveedores=True, log=lambda s: None)
    assert load_workbook(cfg.planilla)["Proveedores"]["B5"].value is None


def test_agrega_hojas_de_areas_nuevas(tmp_path):
    """A workbook from before Persona / ServiciosYa gets them as empty copies of Hogar; the rest is kept."""
    cfg = _cfg(tmp_path)
    preparar.crear_planilla_vacia(cfg.planilla)
    wb = load_workbook(cfg.planilla)
    del wb["Persona"], wb["ServiciosYa"]
    wb["Hogar"]["A12"], wb["Hogar"]["G12"] = "Plomero X", 1000
    wb.save(cfg.planilla)
    assert preparar.hojas_faltantes(cfg.planilla) == ["Persona", "ServiciosYa"]
    assert preparar.plan_setup(cfg)["hojas"] == ["Persona", "ServiciosYa"]

    res = preparar.agregar_hojas(cfg, tmp_path / "backups", log=lambda s: None)
    assert res["hojas"] == ["Persona", "ServiciosYa"] and len(res["respaldos"]) == 2
    wb = load_workbook(cfg.planilla)
    assert wb.sheetnames[:4] == ["Vehicular", "Hogar", "Persona", "ServiciosYa"]
    ws = wb["Persona"]
    assert ws["A1"].value == "Deuda a pagar · Persona" and ws["A11"].value == "Proveedor"
    assert ws["A12"].value is None and ws["H12"].value == '=IF(G12="","",ROUND(G12*$B$4,2))'
    assert [d.formula1 for d in ws.data_validations.dataValidation] == ["=Listas!$C$2:$C$50"]
    assert wb["Hogar"]["A12"].value == "Plomero X"
    assert preparar.hojas_faltantes(cfg.planilla) == []
