"""G2: closing a round (tanda): export, note its batches, optionally empty the workbook; and undo it."""
from openpyxl import load_workbook

from appfacturas import ejecutor, excel, tandas
from test_ejecutor_medias import _armar
from test_lote_sintetico import factura, sandbox  # noqa: F401  (fixture)


def _filas(cfg):
    ws = load_workbook(cfg.planilla)["Vehicular"]
    return [ws[f"A{r}"].value for r in range(excel.PRIMERA, excel.ULTIMA + 1) if ws[f"A{r}"].value]


def test_cerrar_vaciar_y_deshacer(sandbox, tmp_path):
    datos = tmp_path / "datos"
    factura(sandbox.entrada / "Vehicular" / "uno.pdf", 121)
    lote, trabajo = _armar(sandbox, tmp_path)
    ejecutor.Ejecutor(sandbox, lote, trabajo, log=lambda s: None).ejecutar()
    assert tandas.lotes_abiertos(datos, sandbox.planilla) == [trabajo.name]
    assert tandas.siguiente(datos, sandbox.planilla) == 1

    t = tandas.cerrar(sandbox, datos, vaciar=True, log=lambda s: None)
    exportada = sandbox.planilla.parent / "Planillas exportadas" / "DEUDA A PAGAR ARG tanda 1.xlsx"
    assert t["n"] == 1 and t["lotes"] == [trabajo.name] and t["vaciada"]
    assert exportada.is_file() and len(_filas(sandbox)) == 0
    assert len(_filas(type("C", (), {"planilla": exportada})())) == 1        # the copy keeps the rows
    assert not (sandbox.planilla.parent / "Respaldos").exists()               # only the local backup
    assert tandas.lotes_abiertos(datos, sandbox.planilla) == []
    assert tandas.tanda_de(datos) == {trabajo.name: 1}
    assert tandas.siguiente(datos, sandbox.planilla) == 2

    assert tandas.deshacer(sandbox, datos, log=lambda s: None) == []
    assert len(_filas(sandbox)) == 1 and not exportada.exists()
    assert tandas.siguiente(datos, sandbox.planilla) == 1
    assert tandas.lotes_abiertos(datos, sandbox.planilla) == [trabajo.name]


def test_cerrar_sin_vaciar(sandbox, tmp_path):
    datos = tmp_path / "datos"
    t = tandas.cerrar(sandbox, datos, nombre="Pagos de octubre", log=lambda s: None)
    assert t["exportada"].endswith("Pagos de octubre.xlsx") and not t["vaciada"] and t["lotes"] == []
